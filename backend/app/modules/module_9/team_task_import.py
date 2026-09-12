# Bulk XLSX import/export for the Team Task lookup: building the
# downloadable template, applying an uploaded workbook row by row, and
# exporting the full list back out. Mirrors KarTracker's
# kar_status_import.py (same shape, just for a different single-field
# lookup table).
#
# Import is best-effort: every row is validated and applied independently
# inside its own SAVEPOINT (db.begin_nested()), so one bad row rolls back
# only itself instead of aborting the whole batch. A row naming a team
# task that already exists is REJECTED as an error rather than upserted
# — same rule KarTracker uses for its own lookups.

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.models.team_task import TeamTask
from app.schemas.masterdata import TeamTaskImportRowResult

TEMPLATE_COLUMNS = ["team_tasks"]


def build_team_task_template_xlsx() -> bytes:
    """The downloadable XLSX template: just a bold header row. Deliberately
    no example data row — since a re-uploaded, unedited row would silently
    create a real team task (import rejects duplicates rather than
    upserting, so it wouldn't even get caught as "already exists" the
    first time).
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TeamTasks"

    sheet.append(TEMPLATE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: object) -> str | None:
    """Empty/blank cells mean "not set"; anything else is stringified and trimmed."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def import_team_tasks_from_xlsx(db: Session, file_bytes: bytes) -> list[TeamTaskImportRowResult]:
    """Apply every row of an uploaded XLSX workbook's active sheet,
    returning one result per row. Row numbers match the actual
    spreadsheet row (header is row 1, so data starts at row 2) — matching
    what a person sees if they open the file themselves.
    """
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    header_row = next(rows, None)
    if header_row is None:
        return []

    column_index = {str(name).strip().lower(): index for index, name in enumerate(header_row) if name is not None}

    def cell(row: tuple, key: str) -> object:
        index = column_index.get(key)
        return row[index] if index is not None and index < len(row) else None

    results: list[TeamTaskImportRowResult] = []

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(value is None for value in row):
            continue  # a blank trailing row — nothing to report

        team_tasks = _cell_text(cell(row, "team_tasks"))
        savepoint = db.begin_nested()
        try:
            if team_tasks is None:
                raise ValueError("team_tasks is required")

            existing = db.scalar(select(TeamTask).where(TeamTask.team_tasks == team_tasks))
            if existing is not None:
                raise ValueError("A team task with this name already exists")

            db.add(TeamTask(team_tasks=team_tasks))

            savepoint.commit()
            results.append(
                TeamTaskImportRowResult(row_number=row_number, team_tasks=team_tasks, outcome="created", detail=None)
            )
        except (ValueError, IntegrityError) as error:
            savepoint.rollback()
            detail = str(error) if isinstance(error, ValueError) else "This row conflicts with existing data"
            results.append(
                TeamTaskImportRowResult(row_number=row_number, team_tasks=team_tasks, outcome="error", detail=detail)
            )

    db.commit()
    return results


def export_team_tasks_to_xlsx(db: Session) -> bytes:
    """Every team task as a single-sheet workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TeamTasks"
    sheet.append(["id", "team_tasks"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    team_tasks = db.scalars(select(TeamTask).order_by(TeamTask.team_tasks)).all()
    for team_task in team_tasks:
        sheet.append([team_task.id, team_task.team_tasks])

    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 24

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
