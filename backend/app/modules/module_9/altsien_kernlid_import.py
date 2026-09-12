# Bulk XLSX import/export for Altsien Kernleden contacts: building the
# downloadable template, applying an uploaded workbook row by row, and
# exporting the full list back out. Mirrors KarTracker's kar_import.py's
# overall shape, simplified for a flat table with no foreign keys and no
# unique constraint (so, unlike every other Data Upload/Download tile,
# there is no "already exists" rejection here — every valid row is simply
# added as a new contact).
#
# Import is best-effort: every row is validated and applied independently
# inside its own SAVEPOINT (db.begin_nested()), so one bad row rolls back
# only itself instead of aborting the whole batch.

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.models.altsien_kernlid import AltsienKernlid
from app.schemas.masterdata import AltsienKernlidImportRowResult

TEMPLATE_COLUMNS = ["first_name", "name", "telephone_number", "email"]


def build_altsien_kernlid_template_xlsx() -> bytes:
    """The downloadable XLSX template: just a bold header row."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "AltsienKernleden"

    sheet.append(TEMPLATE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: object) -> str | None:
    """Empty/blank cells mean "not set"; anything else is stringified and trimmed."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _telephone_number_cell_text(value: object) -> str | None:
    """Like _cell_text, but rejects a cell Excel stored as a number. A
    phone number typed into an unformatted cell is silently coerced to a
    number by Excel — dropping a leading "0" (e.g. "0471234567" becomes
    471234567) or, for a long number, switching to scientific notation —
    so str(value) would silently import corrupted data instead of what
    the person actually typed. Format the column as Text to avoid this.
    """
    if isinstance(value, (int, float)):
        raise ValueError(
            "telephone_number was stored as a number by Excel, which may have dropped a leading zero — "
            "format the column as Text and re-enter the value"
        )
    return _cell_text(value)


def import_altsien_kernleden_from_xlsx(db: Session, file_bytes: bytes) -> list[AltsienKernlidImportRowResult]:
    """Apply every row of an uploaded XLSX workbook's active sheet,
    returning one result per row. Row numbers match the actual
    spreadsheet row (header is row 1, so data starts at row 2) — matching
    what a person sees if they open the file themselves.
    """
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    header_row = next(rows, None)
    if header_row is None:
        return []

    column_index = {str(name).strip().lower(): index for index, name in enumerate(header_row) if name is not None}

    def cell(row: tuple, key: str) -> object:
        index = column_index.get(key)
        return row[index] if index is not None and index < len(row) else None

    results: list[AltsienKernlidImportRowResult] = []

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(value is None for value in row):
            continue  # a blank trailing row — nothing to report

        name = _cell_text(cell(row, "name"))
        savepoint = db.begin_nested()
        try:
            first_name = _cell_text(cell(row, "first_name"))
            telephone_number = _telephone_number_cell_text(cell(row, "telephone_number"))
            email = _cell_text(cell(row, "email"))

            if first_name is None:
                raise ValueError("first_name is required")
            if name is None:
                raise ValueError("name is required")
            if telephone_number is None:
                raise ValueError("telephone_number is required")
            if email is None:
                raise ValueError("email is required")

            db.add(
                AltsienKernlid(
                    first_name=first_name,
                    name=name,
                    telephone_number=telephone_number,
                    email=email,
                )
            )

            savepoint.commit()
            results.append(
                AltsienKernlidImportRowResult(row_number=row_number, name=name, outcome="created", detail=None)
            )
        except (ValueError, IntegrityError) as error:
            savepoint.rollback()
            detail = str(error) if isinstance(error, ValueError) else "This row conflicts with existing data"
            results.append(
                AltsienKernlidImportRowResult(row_number=row_number, name=name, outcome="error", detail=detail)
            )

    db.commit()
    return results


def export_altsien_kernleden_to_xlsx(db: Session) -> bytes:
    """Every Altsien Kernleden contact as a single-sheet workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "AltsienKernleden"
    sheet.append(["id", *TEMPLATE_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    contacts = db.scalars(select(AltsienKernlid).order_by(AltsienKernlid.name)).all()
    for contact in contacts:
        sheet.append([contact.id, contact.first_name, contact.name, contact.telephone_number, contact.email])

    for index, column in enumerate(["id", *TEMPLATE_COLUMNS], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
