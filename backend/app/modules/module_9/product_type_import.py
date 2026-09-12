# Bulk XLSX import/export for the Product Type lookup: building the
# downloadable template, applying an uploaded workbook row by row, and
# exporting the full list back out. Mirrors KarTracker's
# kar_status_import.py (same shape, just for a different single-field
# lookup table).
#
# Import is best-effort: every row is validated and applied independently
# inside its own SAVEPOINT (db.begin_nested()), so one bad row rolls back
# only itself instead of aborting the whole batch. A row naming a type
# that already exists is REJECTED as an error rather than upserted — same
# rule KarTracker uses for its own lookups.

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.models.product_type import ProductType
from app.schemas.masterdata import ProductTypeImportRowResult

TEMPLATE_COLUMNS = ["name"]


def build_product_type_template_xlsx() -> bytes:
    """The downloadable XLSX template: just a bold header row. Deliberately
    no example data row — since a re-uploaded, unedited row would silently
    create a real type (import rejects duplicates rather than upserting,
    so it wouldn't even get caught as "already exists" the first time).
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Type"

    sheet.append(TEMPLATE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: object) -> str | None:
    """Empty/blank cells mean "not set"; anything else is stringified and trimmed."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def import_product_types_from_xlsx(db: Session, file_bytes: bytes) -> list[ProductTypeImportRowResult]:
    """Apply every row of an uploaded XLSX workbook's active sheet,
    returning one result per row. Row numbers match the actual
    spreadsheet row (header is row 1, so data starts at row 2) — matching
    what a person sees if they open the file themselves.
    """
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    header_row = next(rows, None)
    if header_row is None:
        return []

    column_index = {str(name).strip().lower(): index for index, name in enumerate(header_row) if name is not None}

    def cell(row: tuple, key: str) -> object:
        index = column_index.get(key)
        return row[index] if index is not None and index < len(row) else None

    results: list[ProductTypeImportRowResult] = []

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(value is None for value in row):
            continue  # a blank trailing row — nothing to report

        name = _cell_text(cell(row, "name"))
        savepoint = db.begin_nested()
        try:
            if name is None:
                raise ValueError("name is required")

            existing = db.scalar(select(ProductType).where(ProductType.name == name))
            if existing is not None:
                raise ValueError("A type with this name already exists")

            db.add(ProductType(name=name))

            savepoint.commit()
            results.append(
                ProductTypeImportRowResult(row_number=row_number, name=name, outcome="created", detail=None)
            )
        except (ValueError, IntegrityError) as error:
            savepoint.rollback()
            detail = str(error) if isinstance(error, ValueError) else "This row conflicts with existing data"
            results.append(
                ProductTypeImportRowResult(row_number=row_number, name=name, outcome="error", detail=detail)
            )

    db.commit()
    return results


def export_product_types_to_xlsx(db: Session) -> bytes:
    """Every product type as a single-sheet workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Type"
    sheet.append(["id", "name"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    product_types = db.scalars(select(ProductType).order_by(ProductType.name)).all()
    for product_type in product_types:
        sheet.append([product_type.id, product_type.name])

    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 24

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
