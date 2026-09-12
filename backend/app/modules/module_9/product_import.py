# Bulk XLSX import/export for Products: building the downloadable
# template, applying an uploaded workbook row by row, and exporting the
# full list back out. Mirrors KarTracker's kar_import.py's shape most
# closely of all the MasterData tables — several optional FKs given by
# name, resolved against their lookup tables on import, plus boolean
# checkbox columns.
#
# Import is best-effort: every row is validated and applied independently
# inside its own SAVEPOINT (db.begin_nested()), so one bad row rolls back
# only itself instead of aborting the whole batch. Product.name has no
# unique constraint (matching create_product in router.py, which has no
# IntegrityError handling either) — every valid row is simply inserted,
# there is no "already exists" rejection here.

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.models.product import Product
from app.db.models.product_category import ProductCategory
from app.db.models.product_limit import ProductLimit
from app.db.models.product_type import ProductType
from app.db.models.warehouse import Warehouse
from app.schemas.masterdata import ProductImportRowResult

TEMPLATE_COLUMNS = [
    "name",
    "type_name",
    "warehouse_name",
    "warehouse_location",
    "category_name",
    "is_consumable",
    "is_blocked",
    "is_logistics_product",
    "limit_name",
    "description",
]


def build_product_template_xlsx() -> bytes:
    """The downloadable XLSX template: just a bold header row."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"

    sheet.append(TEMPLATE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: object) -> str | None:
    """Empty/blank cells mean "not set"; anything else is stringified and trimmed."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


_TRUE_CELL_VALUES = {"1", "true", "yes"}
_FALSE_CELL_VALUES = {"0", "false", "no"}


def _parse_bool_cell(column_name: str, value: object) -> bool:
    """Blank means False, matching the model's own default=False. Any
    other value must be a recognized true/false spelling — a typo (e.g.
    "tru") is rejected as an error instead of silently becoming True,
    which is how a permissive "anything non-blank is truthy" rule would
    otherwise treat it.
    """
    text = _cell_text(value)
    if text is None:
        return False
    normalized = text.lower()
    if normalized in _TRUE_CELL_VALUES:
        return True
    if normalized in _FALSE_CELL_VALUES:
        return False
    raise ValueError(f'{column_name} "{text}" is not recognized — use "true"/"false", "1"/"0", or "yes"/"no"')


def _resolve_type_id(db: Session, type_name: str | None) -> int | None:
    if type_name is None:
        return None
    product_type = db.scalar(select(ProductType).where(ProductType.name.ilike(type_name)))
    if product_type is None:
        raise ValueError(f'Type "{type_name}" not found')
    return product_type.id


def _resolve_warehouse_id(db: Session, warehouse_name: str | None) -> int | None:
    if warehouse_name is None:
        return None
    warehouse = db.scalar(select(Warehouse).where(Warehouse.name.ilike(warehouse_name)))
    if warehouse is None:
        raise ValueError(f'Warehouse "{warehouse_name}" not found')
    return warehouse.id


def _resolve_category_id(db: Session, category_name: str | None) -> int | None:
    if category_name is None:
        return None
    category = db.scalar(select(ProductCategory).where(ProductCategory.name.ilike(category_name)))
    if category is None:
        raise ValueError(f'Category "{category_name}" not found')
    return category.id


def _resolve_limit_id(db: Session, limit_name: str | None) -> int | None:
    if limit_name is None:
        return None
    product_limit = db.scalar(select(ProductLimit).where(ProductLimit.name.ilike(limit_name)))
    if product_limit is None:
        raise ValueError(f'Limit "{limit_name}" not found')
    return product_limit.id


def import_products_from_xlsx(db: Session, file_bytes: bytes) -> list[ProductImportRowResult]:
    """Apply every row of an uploaded XLSX workbook's active sheet,
    returning one result per row. Row numbers match the actual
    spreadsheet row (header is row 1, so data starts at row 2) — matching
    what a person sees if they open the file themselves.
    """
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    header_row = next(rows, None)
    if header_row is None:
        return []

    column_index = {str(name).strip().lower(): index for index, name in enumerate(header_row) if name is not None}

    def cell(row: tuple, key: str) -> object:
        index = column_index.get(key)
        return row[index] if index is not None and index < len(row) else None

    results: list[ProductImportRowResult] = []

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(value is None for value in row):
            continue  # a blank trailing row — nothing to report

        name = _cell_text(cell(row, "name"))
        savepoint = db.begin_nested()
        try:
            if name is None:
                raise ValueError("name is required")

            new_product = Product(
                name=name,
                type_id=_resolve_type_id(db, _cell_text(cell(row, "type_name"))),
                warehouse_id=_resolve_warehouse_id(db, _cell_text(cell(row, "warehouse_name"))),
                warehouse_location=_cell_text(cell(row, "warehouse_location")),
                category_id=_resolve_category_id(db, _cell_text(cell(row, "category_name"))),
                is_consumable=_parse_bool_cell("is_consumable", cell(row, "is_consumable")),
                is_blocked=_parse_bool_cell("is_blocked", cell(row, "is_blocked")),
                is_logistics_product=_parse_bool_cell("is_logistics_product", cell(row, "is_logistics_product")),
                limit_id=_resolve_limit_id(db, _cell_text(cell(row, "limit_name"))),
                description=_cell_text(cell(row, "description")),
            )
            db.add(new_product)

            savepoint.commit()
            results.append(ProductImportRowResult(row_number=row_number, name=name, outcome="created", detail=None))
        except (ValueError, IntegrityError) as error:
            savepoint.rollback()
            detail = str(error) if isinstance(error, ValueError) else "This row conflicts with existing data"
            results.append(ProductImportRowResult(row_number=row_number, name=name, outcome="error", detail=detail))

    db.commit()
    return results


def export_products_to_xlsx(db: Session) -> bytes:
    """Every product as a single-sheet workbook, with its FKs resolved
    back to names to match the import template's own columns.
    """
    types = {product_type.id: product_type.name for product_type in db.scalars(select(ProductType)).all()}
    warehouses = {warehouse.id: warehouse.name for warehouse in db.scalars(select(Warehouse)).all()}
    categories = {category.id: category.name for category in db.scalars(select(ProductCategory)).all()}
    limits = {product_limit.id: product_limit.name for product_limit in db.scalars(select(ProductLimit)).all()}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(["id", *TEMPLATE_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    products = db.scalars(select(Product).order_by(Product.name)).all()
    for product in products:
        sheet.append(
            [
                product.id,
                product.name,
                types.get(product.type_id, "") if product.type_id is not None else "",
                warehouses.get(product.warehouse_id, "") if product.warehouse_id is not None else "",
                product.warehouse_location,
                categories.get(product.category_id, "") if product.category_id is not None else "",
                product.is_consumable,
                product.is_blocked,
                product.is_logistics_product,
                limits.get(product.limit_id, "") if product.limit_id is not None else "",
                product.description,
            ]
        )

    for index, column in enumerate(["id", *TEMPLATE_COLUMNS], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
