# Shared pieces of MasterData's bulk XLSX import/export (the "Data
# Upload/Download" screen): reading an uploaded sheet, parsing cells, the
# optional "id" column that turns a row into an update, writing a workbook,
# and the whole import/export for the simple one-name lookup tables
# (warehouses, types, categories, Limieten, team locations, team tasks,
# delivery methods). Each table keeps its own <table>_import.py with the
# functions router.py calls; this file only removes the copy-paste.
#
# The rule every MasterData import follows: every row is applied inside its
# own SAVEPOINT (one bad row rolls back only itself), and the caller gets one
# result per row. A row whose "id" names an existing record UPDATES it, an
# empty "id" CREATES a new one, an unknown "id" is a row error — so an export
# can be edited and uploaded again (its first column is that id).

import io
from collections.abc import Callable, Iterator
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

# What can happen to one uploaded row ("unchanged" is for the link tables:
# a link that already exists is left as is) — re-exported for the importers.
from app.schemas.masterdata import ImportOutcome

# The optional first column of every template and export.
ID_COLUMN = "id"

# Accepted spellings for a yes/no cell, compared lowercase.
TRUE_VALUES = {"true", "yes", "ja", "1", "y"}
FALSE_VALUES = {"false", "no", "nee", "0", "n"}

# Generic message for a database constraint a row runs into.
CONFLICT_DETAIL = "This row conflicts with existing data"


def cell_text(value: object) -> str | None:
    """Empty/blank cells mean "not set"; anything else is stringified and trimmed."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def parse_yes_no(column: str, value: object, *, default: bool) -> bool:
    """A yes/no cell: blank gives the default, a typo is a row error (never
    silently true)."""
    if isinstance(value, bool):
        return value
    text = cell_text(value)
    if text is None:
        return default
    lowered = text.lower()
    if lowered in TRUE_VALUES:
        return True
    if lowered in FALSE_VALUES:
        return False
    raise ValueError(f'{column} "{text}" is not a valid yes/no value')


class UploadedRow:
    """One data row of an uploaded sheet, read by column name."""

    def __init__(self, row_number: int, values: tuple, column_index: dict[str, int]):
        self.row_number = row_number
        self._values = values
        self._column_index = column_index

    def has(self, column: str) -> bool:
        """Whether the uploaded sheet has this column at all."""
        return column in self._column_index

    def value(self, column: str) -> object:
        """The raw cell value (None when the column or cell is missing)."""
        index = self._column_index.get(column)
        return self._values[index] if index is not None and index < len(self._values) else None

    def text(self, column: str) -> str | None:
        """The trimmed text of a cell, None when blank."""
        return cell_text(self.value(column))


def read_rows(file_bytes: bytes) -> Iterator[UploadedRow]:
    """Every non-blank data row of the workbook's active sheet. Row numbers
    match the spreadsheet itself (header is row 1, data starts at row 2);
    header names are matched ignoring case and surrounding spaces.
    """
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    rows = workbook.active.iter_rows(values_only=True)

    header_row = next(rows, None)
    if header_row is None:
        return

    column_index = {str(name).strip().lower(): index for index, name in enumerate(header_row) if name is not None}
    for row_number, values in enumerate(rows, start=2):
        if values is None or all(value is None for value in values):
            continue  # a blank (trailing) row — nothing to report
        yield UploadedRow(row_number, values, column_index)


def find_by_id(db: Session, model: type, row: UploadedRow) -> Any | None:
    """The record named by the row's "id" cell, or None when it's blank (=
    create). A non-number or an id that doesn't exist is a row error."""
    text = row.text(ID_COLUMN)
    if text is None:
        return None
    try:
        record_id = int(float(text))
    except ValueError:
        raise ValueError(f'id "{text}" is not a number') from None
    record = db.get(model, record_id)
    if record is None:
        raise ValueError(f"id {record_id} not found")
    return record


class RowError(ValueError):
    """A row error that remembers the row's display key (e.g. its name)."""

    def __init__(self, message: str, row_key: str | None):
        super().__init__(message)
        self.row_key = row_key


def run_import(
    rows: Iterator[UploadedRow],
    db: Session,
    apply_row: Callable[[UploadedRow], tuple[str | None, ImportOutcome]],
    make_result: Callable[[int, str | None, ImportOutcome, str | None], Any],
) -> list:
    """Apply every row inside its own SAVEPOINT and collect one result per row.

    apply_row returns (the row's display key, outcome) or raises ValueError
    with the reason — a RowError also carries the key, so an error result
    still shows which row (e.g. which name) failed.
    """
    results = []
    for row in rows:
        savepoint = db.begin_nested()
        try:
            key, outcome = apply_row(row)
            savepoint.commit()
            results.append(make_result(row.row_number, key, outcome, None))
        except (ValueError, IntegrityError) as error:
            savepoint.rollback()
            detail = str(error) if isinstance(error, ValueError) else CONFLICT_DETAIL
            results.append(make_result(row.row_number, getattr(error, "row_key", None), "error", detail))
    db.commit()
    return results


def build_xlsx(title: str, header: list[str], rows: list[list] | None = None, *, min_width: int = 18) -> bytes:
    """A single-sheet workbook: a bold header row, then the given rows."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for values in rows or []:
        sheet.append(values)

    # The id column stays narrow; the others fit their header.
    for index, column in enumerate(header, start=1):
        width = 8 if column == ID_COLUMN else max(len(column) + 2, min_width)
        sheet.column_dimensions[get_column_letter(index)].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- The one-name lookup tables --------------------------------------------


def lookup_template_xlsx(title: str, column: str) -> bytes:
    """A lookup's template: only the header ("id" + the name column). No
    example row on purpose: re-uploaded unedited, it would create a real record."""
    return build_xlsx(title, [ID_COLUMN, column], min_width=24)


def lookup_export_xlsx(db: Session, *, model: type, field: str, title: str, column: str) -> bytes:
    """Every record of a lookup table, sorted by name, in the template's columns."""
    records = db.scalars(select(model).order_by(getattr(model, field))).all()
    return build_xlsx(
        title, [ID_COLUMN, column], [[record.id, getattr(record, field)] for record in records], min_width=24
    )


def import_lookup_xlsx(
    db: Session,
    file_bytes: bytes,
    *,
    model: type,
    field: str,
    column: str,
    duplicate_message: str,
    make_result: Callable[[int, str | None, ImportOutcome, str | None], Any],
) -> list:
    """Import a one-name lookup table: create a new name, or rename the
    record named by "id". A name another record already has is a row error
    (the same uniqueness rule as the screen)."""

    def apply_row(row: UploadedRow) -> tuple[str | None, ImportOutcome]:
        name = row.text(column)
        try:
            if name is None:
                raise ValueError(f"{column} is required")
            record = find_by_id(db, model, row)

            # The name must stay unique: another record with it is an error.
            clash = db.scalar(select(model).where(getattr(model, field) == name))
            if clash is not None and clash is not record:
                raise ValueError(duplicate_message)

            if record is None:
                db.add(model(**{field: name}))
                return name, "created"
            setattr(record, field, name)
            return name, "updated"
        except ValueError as error:
            raise RowError(str(error), name) from None

    return run_import(read_rows(file_bytes), db, apply_row, make_result)
