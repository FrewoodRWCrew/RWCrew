# Bulk XLSX import/export for Festivals: building the downloadable
# template, applying an uploaded workbook row by row, and exporting the
# full list back out. Mirrors KarTracker's kar_import.py's shape (an FK
# given by name, resolved against its lookup table on import), but for
# Festival's own fields.
#
# Import is best-effort: every row is validated and applied independently
# inside its own SAVEPOINT (db.begin_nested()), so one bad row rolls back
# only itself instead of aborting the whole batch. Festival.name has no
# unique constraint (matching create_festival in router.py, which has no
# IntegrityError handling either) — every valid row is simply inserted,
# there is no "already exists" rejection here.

import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.models.festival import Festival
from app.db.models.season import Season
from app.schemas.masterdata import FestivalImportRowResult

TEMPLATE_COLUMNS = ["name", "start_date", "end_date", "season_name"]


def build_festival_template_xlsx() -> bytes:
    """The downloadable XLSX template: just a bold header row."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Festivals"

    sheet.append(TEMPLATE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: object) -> str | None:
    """Empty/blank cells mean "not set"; anything else is stringified and trimmed."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_date_cell(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value).strip())


def _resolve_season_id(db: Session, season_name: str | None) -> int:
    if season_name is None:
        raise ValueError("season_name is required")
    season = db.scalar(select(Season).where(Season.name.ilike(season_name)))
    if season is None:
        raise ValueError(f'Season "{season_name}" not found')
    return season.id


def import_festivals_from_xlsx(db: Session, file_bytes: bytes) -> list[FestivalImportRowResult]:
    """Apply every row of an uploaded XLSX workbook's active sheet,
    returning one result per row. Row numbers match the actual
    spreadsheet row (header is row 1, so data starts at row 2) — matching
    what a person sees if they open the file themselves.
    """
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    header_row = next(rows, None)
    if header_row is None:
        return []

    column_index = {str(name).strip().lower(): index for index, name in enumerate(header_row) if name is not None}

    def cell(row: tuple, key: str) -> object:
        index = column_index.get(key)
        return row[index] if index is not None and index < len(row) else None

    results: list[FestivalImportRowResult] = []

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(value is None for value in row):
            continue  # a blank trailing row — nothing to report

        name = _cell_text(cell(row, "name"))
        savepoint = db.begin_nested()
        try:
            if name is None:
                raise ValueError("name is required")

            start_date = _parse_date_cell(cell(row, "start_date"))
            if start_date is None:
                raise ValueError("start_date is required")

            end_date = _parse_date_cell(cell(row, "end_date"))
            if end_date is None:
                raise ValueError("end_date is required")

            season_id = _resolve_season_id(db, _cell_text(cell(row, "season_name")))

            db.add(Festival(name=name, start_date=start_date, end_date=end_date, season_id=season_id))

            savepoint.commit()
            results.append(FestivalImportRowResult(row_number=row_number, name=name, outcome="created", detail=None))
        except (ValueError, IntegrityError) as error:
            savepoint.rollback()
            detail = str(error) if isinstance(error, ValueError) else "This row conflicts with existing data"
            results.append(FestivalImportRowResult(row_number=row_number, name=name, outcome="error", detail=detail))

    db.commit()
    return results


def export_festivals_to_xlsx(db: Session) -> bytes:
    """Every festival as a single-sheet workbook, with its season resolved
    back to a name to match the import template's own columns.
    """
    seasons = {season.id: season.name for season in db.scalars(select(Season)).all()}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Festivals"
    sheet.append(["id", *TEMPLATE_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    festivals = db.scalars(select(Festival).order_by(Festival.name)).all()
    for festival in festivals:
        sheet.append(
            [
                festival.id,
                festival.name,
                festival.start_date,
                festival.end_date,
                seasons.get(festival.season_id, ""),
            ]
        )

    for index, column in enumerate(["id", *TEMPLATE_COLUMNS], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
