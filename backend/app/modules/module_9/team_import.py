# Bulk XLSX import/export for Teams: building the downloadable template,
# applying an uploaded workbook row by row, and exporting the full list
# back out. Mirrors KarTracker's kar_import.py's shape (FKs given by
# name, resolved against their lookup tables on import).
#
# Scope note: this v1 import/export only covers Team's own scalar fields
# (name, location, delivery method, description). The many-to-many links
# to Team Tasks and Altsien Kernleden (see _replace_team_task_links /
# _replace_team_kernlid_links in router.py) are deliberately NOT part of
# this bulk tool — those stay manageable only via the Teams screen itself.
#
# Import is best-effort: every row is validated and applied independently
# inside its own SAVEPOINT (db.begin_nested()), so one bad row rolls back
# only itself instead of aborting the whole batch. Team.name IS unique
# (matching create_team in router.py), so a row naming a team that
# already exists is REJECTED as an error rather than upserted.

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.models.delivery_method import DeliveryMethod
from app.db.models.team import Team
from app.db.models.team_location import TeamLocation
from app.schemas.masterdata import TeamImportRowResult

TEMPLATE_COLUMNS = ["name", "location_name", "delivery_method_name", "description"]


def build_team_template_xlsx() -> bytes:
    """The downloadable XLSX template: just a bold header row. Deliberately
    no example data row — since a re-uploaded, unedited row would silently
    create a real team (import rejects duplicates rather than upserting,
    so it wouldn't even get caught as "already exists" the first time).
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Teams"

    sheet.append(TEMPLATE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: object) -> str | None:
    """Empty/blank cells mean "not set"; anything else is stringified and trimmed."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _resolve_location_id(db: Session, location_name: str | None) -> int | None:
    if location_name is None:
        return None
    location = db.scalar(select(TeamLocation).where(TeamLocation.location.ilike(location_name)))
    if location is None:
        raise ValueError(f'Team location "{location_name}" not found')
    return location.id


def _resolve_delivery_method_id(db: Session, delivery_method_name: str | None) -> int | None:
    if delivery_method_name is None:
        return None
    delivery_method = db.scalar(
        select(DeliveryMethod).where(DeliveryMethod.delivery_method.ilike(delivery_method_name))
    )
    if delivery_method is None:
        raise ValueError(f'Delivery method "{delivery_method_name}" not found')
    return delivery_method.id


def import_teams_from_xlsx(db: Session, file_bytes: bytes) -> list[TeamImportRowResult]:
    """Apply every row of an uploaded XLSX workbook's active sheet,
    returning one result per row. Row numbers match the actual
    spreadsheet row (header is row 1, so data starts at row 2) — matching
    what a person sees if they open the file themselves.
    """
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    header_row = next(rows, None)
    if header_row is None:
        return []

    column_index = {str(name).strip().lower(): index for index, name in enumerate(header_row) if name is not None}

    def cell(row: tuple, key: str) -> object:
        index = column_index.get(key)
        return row[index] if index is not None and index < len(row) else None

    results: list[TeamImportRowResult] = []

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(value is None for value in row):
            continue  # a blank trailing row — nothing to report

        name = _cell_text(cell(row, "name"))
        savepoint = db.begin_nested()
        try:
            if name is None:
                raise ValueError("name is required")

            existing = db.scalar(select(Team).where(Team.name == name))
            if existing is not None:
                raise ValueError("A team with this name already exists")

            new_team = Team(
                name=name,
                location_id=_resolve_location_id(db, _cell_text(cell(row, "location_name"))),
                delivery_method_id=_resolve_delivery_method_id(db, _cell_text(cell(row, "delivery_method_name"))),
                description=_cell_text(cell(row, "description")),
            )
            db.add(new_team)

            savepoint.commit()
            results.append(TeamImportRowResult(row_number=row_number, name=name, outcome="created", detail=None))
        except (ValueError, IntegrityError) as error:
            savepoint.rollback()
            detail = str(error) if isinstance(error, ValueError) else "This row conflicts with existing data"
            results.append(TeamImportRowResult(row_number=row_number, name=name, outcome="error", detail=detail))

    db.commit()
    return results


def export_teams_to_xlsx(db: Session) -> bytes:
    """Every team's scalar fields as a single-sheet workbook, with its
    location/delivery method resolved back to names to match the import
    template's own columns. Task/kernlid links are not exported — see the
    scope note at the top of this file.
    """
    locations = {location.id: location.location for location in db.scalars(select(TeamLocation)).all()}
    delivery_methods = {method.id: method.delivery_method for method in db.scalars(select(DeliveryMethod)).all()}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Teams"
    sheet.append(["id", *TEMPLATE_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    teams = db.scalars(select(Team).order_by(Team.name)).all()
    for team in teams:
        sheet.append(
            [
                team.id,
                team.name,
                locations.get(team.location_id, "") if team.location_id is not None else "",
                delivery_methods.get(team.delivery_method_id, "") if team.delivery_method_id is not None else "",
                team.description,
            ]
        )

    for index, column in enumerate(["id", *TEMPLATE_COLUMNS], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
