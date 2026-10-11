# Inboeken from Excel: the downloadable template and reading a filled-in
# copy back. Reading never books anything — it only returns the basket
# lines (product + quantity) for the booking panel, where the person checks
# them and confirms the normal way. So all booking rules stay in
# stock_service.py and a wrong file costs nothing.
#
# Row problems come back as "code|param" (like StockError), translated by
# the frontend under stockMaster.booking.importErrors.<code>.

import io
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models.product import Product
from app.db.models.warehouse import Warehouse
from app.modules.module_4.stock_queries import _bin_label, _names
from app.schemas.stockmaster import BookInImportError, BookInImportLine, BookInImportResponse

TEMPLATE_COLUMNS = ["product_id", "product_name", "bin", "quantity"]
COLUMN_WIDTHS = {"product_id": 12, "product_name": 40, "bin": 24, "quantity": 12}


class BookInImportFileError(Exception):
    """The upload isn't a readable Excel workbook."""


def build_book_in_template_xlsx(db: Session) -> bytes:
    """The template: every product that may be booked in (not blocked), with
    its bin, and an empty quantity column — the person only types quantities.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inboeken"

    # Bold header row that stays in view while scrolling.
    sheet.append(TEMPLATE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    # One row per bookable product, sorted by name like the stock overview.
    warehouses = _names(db, Warehouse, Warehouse.name)
    products = db.scalars(select(Product).where(Product.is_blocked.is_(False)).order_by(Product.name)).all()
    for product in products:
        warehouse_name = warehouses.get(product.warehouse_id) if product.warehouse_id else None
        sheet.append([product.id, product.name, _bin_label(warehouse_name, product.warehouse_location), None])

    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = COLUMN_WIDTHS[column]

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: object) -> str | None:
    """Blank cells mean "not filled in"; anything else is trimmed text."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _whole_number(value: object) -> int | None:
    """A cell as a whole number (Excel gives 5 as 5 or 5.0), None if it isn't one."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    text = _cell_text(value)
    if text is None:
        return None
    try:
        number = float(text.replace(",", "."))
    except ValueError:
        return None
    return int(number) if number.is_integer() else None


def parse_book_in_xlsx(db: Session, file_bytes: bytes) -> BookInImportResponse:
    """Read a filled-in template into basket lines plus per-row problems.
    Row numbers are the spreadsheet's own (header = row 1).
    """
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as error:
        raise BookInImportFileError() from error
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    header_row = next(rows, None)
    if header_row is None:
        return BookInImportResponse(lines=[], errors=[])
    column_index = {str(name).strip().lower(): index for index, name in enumerate(header_row) if name is not None}

    def cell(row: tuple, key: str) -> object:
        index = column_index.get(key)
        return row[index] if index is not None and index < len(row) else None

    # Every product once, to look rows up by id or (unique) name.
    products = db.scalars(select(Product)).all()
    by_id = {product.id: product for product in products}
    by_name: dict[str, list[Product]] = {}
    for product in products:
        by_name.setdefault(product.name.strip().lower(), []).append(product)

    quantities: dict[int, int] = {}  # product id → summed quantity, in file order
    errors: list[BookInImportError] = []

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(value is None for value in row):
            continue

        # No quantity (or 0) = this product isn't delivered: skip silently.
        raw_quantity = cell(row, "quantity")
        if _cell_text(raw_quantity) is None:
            continue
        quantity = _whole_number(raw_quantity)
        name = _cell_text(cell(row, "product_name"))
        label = name or _cell_text(cell(row, "product_id")) or ""
        if quantity is None or quantity < 0:
            errors.append(BookInImportError(row_number=row_number, detail=f"bad_quantity|{_cell_text(raw_quantity)}"))
            continue
        if quantity == 0:
            continue

        # The product: by id first, otherwise by name (only when exactly one product has it).
        product_id = _whole_number(cell(row, "product_id"))
        product: Product | None = None
        if product_id is not None:
            product = by_id.get(product_id)
        elif name is not None:
            matches = by_name.get(name.lower(), [])
            if len(matches) > 1:
                errors.append(BookInImportError(row_number=row_number, detail=f"ambiguous_product|{name}"))
                continue
            product = matches[0] if matches else None
        if product is None:
            errors.append(BookInImportError(row_number=row_number, detail=f"unknown_product|{label}"))
            continue

        # The same product twice in the file adds up.
        quantities[product.id] = quantities.get(product.id, 0) + quantity

    return BookInImportResponse(
        lines=[BookInImportLine(product_id=product_id, quantity=quantity) for product_id, quantity in quantities.items()],
        errors=errors,
    )
