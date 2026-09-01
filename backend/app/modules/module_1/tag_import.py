# Bulk XLSX import for TagManagement: building the downloadable template
# and applying an uploaded workbook row by row. Kept free of FastAPI/
# routing concerns (mirrors file_browser.py), so it's easy to unit test
# on its own.
#
# Import is best-effort: every row is validated and applied independently
# inside its own SAVEPOINT (db.begin_nested()), so one bad row rolls back
# only itself instead of aborting the whole batch — the caller always gets
# back one result per row (created/updated/error), never a partial
# exception. A row whose epc_uid already exists is UPDATED (upsert),
# rather than rejected as a duplicate, so re-exporting/editing/
# re-uploading the same tags is the expected workflow.

import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.models.product import Product
from app.db.models.rfid_tag import RfidTag
from app.schemas.tagscan import RfidTagImportRowResult

# The workbook's columns, in order — also the header row of the
# downloadable template. "product_name" (not an id) since this is filled
# in by a person editing a spreadsheet, matched against Product.name on
# import.
TEMPLATE_COLUMNS = [
    "epc_uid",
    "status",
    "product_name",
    "assigned_serial_number",
    "date_assigned",
    "last_read_at",
    "last_reader_id",
    "last_location",
    "manufacturer",
    "batch_number",
    "notes_1",
    "notes_2",
    "notes_3",
    "notes_4",
    "notes_5",
]

ALLOWED_STATUSES = {"active", "inactive", "lost", "damaged", "retired"}

# One realistic example row, so the template shows the expected format —
# the two date columns are real date/datetime values, not text, so Excel
# displays and lets the user edit them as actual dates.
_EXAMPLE_ROW = {
    "epc_uid": "E200001122334455",
    "status": "active",
    "product_name": "KBC Lint",
    "assigned_serial_number": "SN-001",
    "date_assigned": date(2026, 1, 15),
    "last_read_at": datetime(2026, 2, 1, 10, 30),
    "last_reader_id": "Reader-1",
    "last_location": "Warehouse A",
    "manufacturer": "Impinj",
    "batch_number": "B-42",
    "notes_1": "",
    "notes_2": "",
    "notes_3": "",
    "notes_4": "",
    "notes_5": "",
}


def build_template_xlsx() -> bytes:
    """The downloadable XLSX template: a bold header row plus one example row."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tags"

    sheet.append(TEMPLATE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.append([_EXAMPLE_ROW[column] for column in TEMPLATE_COLUMNS])

    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 14)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: object) -> str | None:
    """Empty/blank cells mean "not set"; anything else is stringified and trimmed."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_date_cell(value: object) -> date | None:
    """A date-formatted Excel cell comes back from openpyxl as an actual
    date/datetime object; a plain-text cell comes back as a string that
    still needs parsing. datetime must be checked before date, since
    datetime is itself a subclass of date.
    """
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value).strip())


def _parse_datetime_cell(value: object) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    return datetime.fromisoformat(str(value).strip())


def _resolve_product_id(db: Session, product_name: str | None) -> int | None:
    if product_name is None:
        return None
    product = db.scalar(select(Product).where(Product.name.ilike(product_name)))
    if product is None:
        raise ValueError(f'Product "{product_name}" not found')
    return product.id


def import_tags_from_xlsx(db: Session, file_bytes: bytes) -> list[RfidTagImportRowResult]:
    """Apply every row of an uploaded XLSX workbook's active sheet,
    returning one result per row. Row numbers match the actual
    spreadsheet row (header is row 1, so data starts at row 2) —
    matching what a person sees if they open the file themselves.
    """
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    header_row = next(rows, None)
    if header_row is None:
        return []

    column_index = {str(name).strip().lower(): index for index, name in enumerate(header_row) if name is not None}

    def cell(row: tuple, key: str) -> object:
        index = column_index.get(key)
        return row[index] if index is not None and index < len(row) else None

    results: list[RfidTagImportRowResult] = []

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(value is None for value in row):
            continue  # a blank trailing row — nothing to report

        epc_uid = _cell_text(cell(row, "epc_uid"))
        savepoint = db.begin_nested()
        try:
            if epc_uid is None:
                raise ValueError("epc_uid is required")

            status_value = (_cell_text(cell(row, "status")) or "active").lower()
            if status_value not in ALLOWED_STATUSES:
                raise ValueError(f'Invalid status "{status_value}"')

            fields = {
                "status": status_value,
                "assigned_product_id": _resolve_product_id(db, _cell_text(cell(row, "product_name"))),
                "assigned_serial_number": _cell_text(cell(row, "assigned_serial_number")),
                "date_assigned": _parse_date_cell(cell(row, "date_assigned")),
                "last_read_at": _parse_datetime_cell(cell(row, "last_read_at")),
                "last_reader_id": _cell_text(cell(row, "last_reader_id")),
                "last_location": _cell_text(cell(row, "last_location")),
                "manufacturer": _cell_text(cell(row, "manufacturer")),
                "batch_number": _cell_text(cell(row, "batch_number")),
                "notes_1": _cell_text(cell(row, "notes_1")),
                "notes_2": _cell_text(cell(row, "notes_2")),
                "notes_3": _cell_text(cell(row, "notes_3")),
                "notes_4": _cell_text(cell(row, "notes_4")),
                "notes_5": _cell_text(cell(row, "notes_5")),
            }

            existing = db.scalar(select(RfidTag).where(RfidTag.epc_uid == epc_uid))
            if existing is not None:
                for field, value in fields.items():
                    setattr(existing, field, value)
                outcome = "updated"
            else:
                db.add(RfidTag(epc_uid=epc_uid, **fields))
                outcome = "created"

            savepoint.commit()
            results.append(RfidTagImportRowResult(row_number=row_number, epc_uid=epc_uid, outcome=outcome, detail=None))
        except (ValueError, IntegrityError) as error:
            savepoint.rollback()
            detail = str(error) if isinstance(error, ValueError) else "This row conflicts with existing data"
            results.append(
                RfidTagImportRowResult(row_number=row_number, epc_uid=epc_uid, outcome="error", detail=detail)
            )

    db.commit()
    return results
