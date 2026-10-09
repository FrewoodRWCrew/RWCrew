# Bulk XLSX import/export for the Scanners registry (used by TagScan's
# "Data Upload/Download" screen): building the downloadable template,
# applying an uploaded workbook row by row, and exporting every scanner in
# that same column layout. Mirrors tag_import.py and is likewise kept free
# of FastAPI/routing concerns.
#
# Import is best-effort: every row is applied inside its own SAVEPOINT, so
# one bad row rolls back only itself — the caller always gets one result
# per row (created/updated/error). A row naming a scanner that already
# exists UPDATES it (upsert), so export → edit → re-upload works. A
# scanner's device API key is never part of the workbook and never touched.

import io
from typing import get_args

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.models.product_type import ProductType
from app.db.models.scanner import Scanner
from app.schemas.tagscan import ScannerImportRowResult, ScannerTechnology

# The workbook's columns, in order — also the header row of the template.
# "type_name" (not an id) since a person fills this in, matched against
# MasterData's ProductType.name on import.
TEMPLATE_COLUMNS = [
    "scanner",
    "type_name",
    "technology",
    "location",
    "description",
    "info1",
    "info2",
    "info3",
]

# The allowed technology values, looked up case-insensitively so "raspberry
# pi 4" in a spreadsheet is stored in its canonical spelling.
_TECHNOLOGIES_BY_LOWER = {value.lower(): value for value in get_args(ScannerTechnology)}

# One realistic example row, so the template shows the expected format.
_EXAMPLE_ROW = {
    "scanner": "Scan_01",
    "type_name": "Reader Hardware",
    "technology": "Raspberry Pi 4",
    "location": "Warehouse A",
    "description": "Gate reader at the loading dock",
    "info1": "",
    "info2": "",
    "info3": "",
}


def _write_header(sheet) -> None:
    """The bold header row shared by the template and the export."""
    sheet.append(TEMPLATE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)


def _to_bytes(workbook: Workbook) -> bytes:
    """Set readable column widths and serialise the workbook."""
    sheet = workbook.active
    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_scanner_template_xlsx() -> bytes:
    """The downloadable XLSX template: a bold header row plus one example row."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Scanners"
    _write_header(sheet)
    sheet.append([_EXAMPLE_ROW[column] for column in TEMPLATE_COLUMNS])
    return _to_bytes(workbook)


def export_scanners_to_xlsx(db: Session) -> bytes:
    """Every scanner as a single-sheet workbook, in the template's columns."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Scanners"
    _write_header(sheet)

    # Type names by id, so each row shows the name a person can edit.
    type_names = {product_type.id: product_type.name for product_type in db.scalars(select(ProductType)).all()}

    for scanner in db.scalars(select(Scanner).order_by(Scanner.scanner)).all():
        sheet.append(
            [
                scanner.scanner,
                type_names.get(scanner.type_id),
                scanner.technology,
                scanner.location,
                scanner.description,
                scanner.info1,
                scanner.info2,
                scanner.info3,
            ]
        )
    return _to_bytes(workbook)


def _cell_text(value: object) -> str | None:
    """Empty/blank cells mean "not set"; anything else is stringified and trimmed."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _resolve_type_id(db: Session, type_name: str | None) -> int:
    """The ProductType id for a name (case-insensitive) — required."""
    if type_name is None:
        raise ValueError("type_name is required")
    product_type = db.scalar(select(ProductType).where(ProductType.name.ilike(type_name)))
    if product_type is None:
        raise ValueError(f'Type "{type_name}" not found')
    return product_type.id


def _resolve_technology(value: str | None) -> str:
    """One of the fixed technology values, in its canonical spelling — required."""
    if value is None:
        raise ValueError("technology is required")
    technology = _TECHNOLOGIES_BY_LOWER.get(value.lower())
    if technology is None:
        raise ValueError(f'Invalid technology "{value}" (allowed: {", ".join(get_args(ScannerTechnology))})')
    return technology


def import_scanners_from_xlsx(db: Session, file_bytes: bytes) -> list[ScannerImportRowResult]:
    """Apply every row of an uploaded XLSX workbook's active sheet,
    returning one result per row. Row numbers match the spreadsheet itself
    (header is row 1, so data starts at row 2).
    """
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    header_row = next(rows, None)
    if header_row is None:
        return []

    column_index = {str(name).strip().lower(): index for index, name in enumerate(header_row) if name is not None}

    def cell(row: tuple, key: str) -> object:
        index = column_index.get(key)
        return row[index] if index is not None and index < len(row) else None

    results: list[ScannerImportRowResult] = []

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(value is None for value in row):
            continue  # a blank trailing row — nothing to report

        # Trimmed, like the Scanners screen does, so it matches CSV lines.
        scanner_name = _cell_text(cell(row, "scanner"))
        savepoint = db.begin_nested()
        try:
            if scanner_name is None:
                raise ValueError("scanner is required")
            if len(scanner_name) > 255:
                raise ValueError("scanner is longer than 255 characters")

            fields = {
                "type_id": _resolve_type_id(db, _cell_text(cell(row, "type_name"))),
                "technology": _resolve_technology(_cell_text(cell(row, "technology"))),
                "location": _cell_text(cell(row, "location")),
                "description": _cell_text(cell(row, "description")),
                "info1": _cell_text(cell(row, "info1")),
                "info2": _cell_text(cell(row, "info2")),
                "info3": _cell_text(cell(row, "info3")),
            }

            # An existing scanner (same name) is updated; its API key stays.
            existing = db.scalar(select(Scanner).where(Scanner.scanner == scanner_name))
            if existing is not None:
                for field, value in fields.items():
                    setattr(existing, field, value)
                outcome = "updated"
            else:
                db.add(Scanner(scanner=scanner_name, **fields))
                outcome = "created"

            savepoint.commit()
            results.append(
                ScannerImportRowResult(row_number=row_number, scanner=scanner_name, outcome=outcome, detail=None)
            )
        except (ValueError, IntegrityError) as error:
            savepoint.rollback()
            detail = str(error) if isinstance(error, ValueError) else "This row conflicts with existing data"
            results.append(
                ScannerImportRowResult(row_number=row_number, scanner=scanner_name, outcome="error", detail=detail)
            )

    db.commit()
    return results
