# Bulk XLSX import/export for the Distributiepunten master data: building
# the downloadable template, applying an uploaded workbook row by row, and
# exporting the full dataset back out. Mirrors kar_import.py's own shape.
#
# Import is best-effort: every row is validated and applied independently
# inside its own SAVEPOINT (db.begin_nested()), so one bad row rolls back
# only itself instead of aborting the whole batch. A row whose name
# already exists is REJECTED as an error rather than upserted — same rule
# as the Karren import.

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.models.kartracker_distributiepunt import KarTrackerDistributiepunt
from app.db.models.user import User
from app.schemas.kartracker import DistributiepuntImportRowResult

# The workbook's columns, in order — also the header row of the
# downloadable template. Altsien Kernlid is given by the user's email
# (not id), since this is filled in by a person editing a spreadsheet,
# matched against the users flagged Altsien Kernlid on import.
TEMPLATE_COLUMNS = [
    "name",
    "latitude",
    "longitude",
    "terrein_positie",
    "altsien_kernlid_email",
]


def build_distributiepunt_template_xlsx() -> bytes:
    """The downloadable XLSX template: just a bold header row. Deliberately
    no example data row — since a re-uploaded, unedited row would silently
    create a real distribution point (import rejects duplicates rather
    than upserting, so it wouldn't even get caught as "already exists"
    the first time).
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Distributiepunten"

    sheet.append(TEMPLATE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: object) -> str | None:
    """Empty/blank cells mean "not set"; anything else is stringified and trimmed."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _resolve_altsien_kernlid_id(db: Session, email: str | None) -> int | None:
    """A blank cell means "no kernlid assigned". An email that doesn't match
    a user flagged Altsien Kernlid is an error rather than silently ignored.
    """
    if email is None:
        return None
    user = db.scalar(select(User).where(func.lower(User.email) == email.lower(), User.is_altsien_kernlid.is_(True)))
    if user is None:
        raise ValueError(f'Altsien Kernlid "{email}" not found')
    return user.id


def import_distributiepunten_from_xlsx(db: Session, file_bytes: bytes) -> list[DistributiepuntImportRowResult]:
    """Apply every row of an uploaded XLSX workbook's active sheet,
    returning one result per row. Row numbers match the actual
    spreadsheet row (header is row 1, so data starts at row 2) — matching
    what a person sees if they open the file themselves.
    """
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    header_row = next(rows, None)
    if header_row is None:
        return []

    column_index = {str(name).strip().lower(): index for index, name in enumerate(header_row) if name is not None}

    def cell(row: tuple, key: str) -> object:
        index = column_index.get(key)
        return row[index] if index is not None and index < len(row) else None

    results: list[DistributiepuntImportRowResult] = []

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(value is None for value in row):
            continue  # a blank trailing row — nothing to report

        name = _cell_text(cell(row, "name"))
        savepoint = db.begin_nested()
        try:
            if name is None:
                raise ValueError("name is required")

            existing = db.scalar(select(KarTrackerDistributiepunt).where(KarTrackerDistributiepunt.name == name))
            if existing is not None:
                raise ValueError("A distribution point with this name already exists")

            new_distributiepunt = KarTrackerDistributiepunt(
                name=name,
                latitude=float(cell(row, "latitude")) if _cell_text(cell(row, "latitude")) else None,
                longitude=float(cell(row, "longitude")) if _cell_text(cell(row, "longitude")) else None,
                terrein_positie=_cell_text(cell(row, "terrein_positie")),
                altsien_kernlid_id=_resolve_altsien_kernlid_id(db, _cell_text(cell(row, "altsien_kernlid_email"))),
            )
            db.add(new_distributiepunt)

            savepoint.commit()
            results.append(
                DistributiepuntImportRowResult(row_number=row_number, name=name, outcome="created", detail=None)
            )
        except (ValueError, IntegrityError) as error:
            savepoint.rollback()
            detail = str(error) if isinstance(error, ValueError) else "This row conflicts with existing data"
            results.append(
                DistributiepuntImportRowResult(row_number=row_number, name=name, outcome="error", detail=detail)
            )

    db.commit()
    return results


def export_distributiepunten_to_xlsx(db: Session) -> bytes:
    """The full Distributiepunten dataset as an XLSX workbook, same
    columns as the import template, with the Altsien Kernlid resolved
    back to the user's email.
    """
    kernlid_emails = {user.id: user.email for user in db.scalars(select(User)).all()}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Distributiepunten"
    sheet.append(TEMPLATE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    distributiepunten = db.scalars(
        select(KarTrackerDistributiepunt).order_by(KarTrackerDistributiepunt.name)
    ).all()
    for distributiepunt in distributiepunten:
        sheet.append(
            [
                distributiepunt.name,
                distributiepunt.latitude,
                distributiepunt.longitude,
                distributiepunt.terrein_positie,
                kernlid_emails.get(distributiepunt.altsien_kernlid_id),
            ]
        )
    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
