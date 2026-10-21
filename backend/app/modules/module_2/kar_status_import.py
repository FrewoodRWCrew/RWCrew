# Bulk XLSX import/export for the KarStatussen lookup: building the
# downloadable template, applying an uploaded workbook row by row, and
# exporting the full list back out. Mirrors kar_import.py's own shape
# (which does the same thing for the fleet registry itself), just for a
# much simpler, single-field lookup.
#
# Import is best-effort: every row is validated and applied independently
# inside its own SAVEPOINT (db.begin_nested()), so one bad row rolls back
# only itself instead of aborting the whole batch. A row naming a status
# that already exists (ignoring case) is reported as "unchanged" rather than
# an error: a status is just a name, so there is nothing to update — this
# keeps an exported list safe to upload again.

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.models.kartracker_kar_status import KarTrackerKarStatus
from app.schemas.kartracker import KarStatusImportRowResult

TEMPLATE_COLUMNS = ["name"]


def build_kar_status_template_xlsx() -> bytes:
    """The downloadable XLSX template: just a bold header row. Deliberately
    no example data row — a re-uploaded, unedited example row would
    silently create a real kar status.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "KarStatussen"

    sheet.append(TEMPLATE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: object) -> str | None:
    """Empty/blank cells mean "not set"; anything else is stringified and trimmed."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def import_kar_statuses_from_xlsx(db: Session, file_bytes: bytes) -> list[KarStatusImportRowResult]:
    """Apply every row of an uploaded XLSX workbook's active sheet,
    returning one result per row. Row numbers match the actual
    spreadsheet row (header is row 1, so data starts at row 2) — matching
    what a person sees if they open the file themselves.
    """
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    header_row = next(rows, None)
    if header_row is None:
        return []

    column_index = {str(name).strip().lower(): index for index, name in enumerate(header_row) if name is not None}

    def cell(row: tuple, key: str) -> object:
        index = column_index.get(key)
        return row[index] if index is not None and index < len(row) else None

    results: list[KarStatusImportRowResult] = []

    # Names already handled in this file (lowercased), so a repeated one is
    # reported instead of silently counted twice.
    seen_names: set[str] = set()

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(value is None for value in row):
            continue  # a blank trailing row — nothing to report

        name = _cell_text(cell(row, "name"))
        savepoint = db.begin_nested()
        try:
            if name is None:
                raise ValueError("name is required")
            if name.lower() in seen_names:
                raise ValueError("This kar status appears more than once in the file")
            seen_names.add(name.lower())

            # Matched ignoring case, so "inhoud ok" doesn't become a second "Inhoud OK".
            existing = db.scalar(select(KarTrackerKarStatus).where(func.lower(KarTrackerKarStatus.name) == name.lower()))
            if existing is None:
                db.add(KarTrackerKarStatus(name=name))
                outcome = "created"
            else:
                outcome = "unchanged"

            savepoint.commit()
            results.append(KarStatusImportRowResult(row_number=row_number, name=name, outcome=outcome, detail=None))
        except (ValueError, IntegrityError) as error:
            savepoint.rollback()
            detail = str(error) if isinstance(error, ValueError) else "This row conflicts with existing data"
            results.append(
                KarStatusImportRowResult(row_number=row_number, name=name, outcome="error", detail=detail)
            )

    db.commit()
    return results


def export_kar_statuses_to_xlsx(db: Session) -> bytes:
    """Every kar status as a single-sheet workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "KarStatussen"
    sheet.append(["id", "name"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    kar_statuses = db.scalars(select(KarTrackerKarStatus).order_by(KarTrackerKarStatus.name)).all()
    for kar_status in kar_statuses:
        sheet.append([kar_status.id, kar_status.name])

    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 24

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
