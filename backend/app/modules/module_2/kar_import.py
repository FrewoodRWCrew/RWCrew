# Bulk XLSX import/export for KarManagement's fleet registry: building the
# downloadable template, applying an uploaded workbook row by row, and
# exporting the full dataset back out. Kept free of FastAPI/routing
# concerns (mirrors module_1/tag_import.py), so it's easy to unit test on
# its own.
#
# Import is best-effort: every row is validated and applied independently
# inside its own SAVEPOINT (db.begin_nested()), so one bad row rolls back
# only itself instead of aborting the whole batch — the caller always gets
# back one result per row (created/error), never a partial exception.
# Unlike TagScan's tag import, a row whose kar_nummer already exists is
# REJECTED as an error rather than upserted — this screen is for
# registering brand-new kars only, editing existing ones stays on the
# KarManagement screen itself.

import io
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.models.kartracker_kar import KarTrackerKar
from app.db.models.kartracker_kar_status import KarTrackerKarStatus
from app.db.models.product import Product
from app.db.models.team import Team
from app.schemas.kartracker import KarImportRowResult

# The workbook's columns, in order — also the header row of the
# downloadable template. Status/team/transport-type are given by name (not
# id), since this is filled in by a person editing a spreadsheet, matched
# against their respective lookup tables on import.
TEMPLATE_COLUMNS = [
    "kar_nummer",
    "status_name",
    "team_name",
    "transport_type_name",
    "last_latitude",
    "last_longitude",
    "last_recorded_at",
]

# One realistic example row, so the template shows the expected format.
_EXAMPLE_ROW = {
    "kar_nummer": "B001",
    "status_name": "Actief",
    "team_name": "",
    "transport_type_name": "KBC Lint",
    "last_latitude": "",
    "last_longitude": "",
    "last_recorded_at": "",
}


def build_kar_template_xlsx() -> bytes:
    """The downloadable XLSX template: a bold header row plus one example row."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Karren"

    sheet.append(TEMPLATE_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.append([_EXAMPLE_ROW[column] for column in TEMPLATE_COLUMNS])

    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: object) -> str | None:
    """Empty/blank cells mean "not set"; anything else is stringified and trimmed."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_datetime_cell(value: object) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    return datetime.fromisoformat(str(value).strip())


def _resolve_status_id(db: Session, status_name: str | None) -> int:
    if status_name is None:
        raise ValueError("status_name is required")
    kar_status = db.scalar(select(KarTrackerKarStatus).where(KarTrackerKarStatus.name.ilike(status_name)))
    if kar_status is None:
        raise ValueError(f'Kar status "{status_name}" not found')
    return kar_status.id


def _resolve_transport_type_id(db: Session, transport_type_name: str | None) -> int:
    if transport_type_name is None:
        raise ValueError("transport_type_name is required")
    product = db.scalar(select(Product).where(Product.name.ilike(transport_type_name)))
    if product is None:
        raise ValueError(f'Transport type "{transport_type_name}" not found')
    return product.id


def _resolve_team_id(db: Session, team_name: str | None) -> int | None:
    if team_name is None:
        return None
    team = db.scalar(select(Team).where(Team.name.ilike(team_name)))
    if team is None:
        raise ValueError(f'Team "{team_name}" not found')
    return team.id


def import_karren_from_xlsx(db: Session, file_bytes: bytes) -> list[KarImportRowResult]:
    """Apply every row of an uploaded XLSX workbook's active sheet,
    returning one result per row. Row numbers match the actual
    spreadsheet row (header is row 1, so data starts at row 2) — matching
    what a person sees if they open the file themselves.
    """
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    header_row = next(rows, None)
    if header_row is None:
        return []

    column_index = {str(name).strip().lower(): index for index, name in enumerate(header_row) if name is not None}

    def cell(row: tuple, key: str) -> object:
        index = column_index.get(key)
        return row[index] if index is not None and index < len(row) else None

    results: list[KarImportRowResult] = []

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(value is None for value in row):
            continue  # a blank trailing row — nothing to report

        kar_nummer = _cell_text(cell(row, "kar_nummer"))
        savepoint = db.begin_nested()
        try:
            if kar_nummer is None:
                raise ValueError("kar_nummer is required")

            existing = db.scalar(select(KarTrackerKar).where(KarTrackerKar.kar_nummer == kar_nummer))
            if existing is not None:
                raise ValueError("A kar with this kar number already exists")

            new_kar = KarTrackerKar(
                kar_nummer=kar_nummer,
                status_id=_resolve_status_id(db, _cell_text(cell(row, "status_name"))),
                team_id=_resolve_team_id(db, _cell_text(cell(row, "team_name"))),
                transport_type_id=_resolve_transport_type_id(db, _cell_text(cell(row, "transport_type_name"))),
                last_latitude=float(cell(row, "last_latitude")) if _cell_text(cell(row, "last_latitude")) else None,
                last_longitude=float(cell(row, "last_longitude")) if _cell_text(cell(row, "last_longitude")) else None,
                last_recorded_at=_parse_datetime_cell(cell(row, "last_recorded_at")),
            )
            db.add(new_kar)

            savepoint.commit()
            results.append(KarImportRowResult(row_number=row_number, kar_nummer=kar_nummer, outcome="created", detail=None))
        except (ValueError, IntegrityError) as error:
            savepoint.rollback()
            detail = str(error) if isinstance(error, ValueError) else "This row conflicts with existing data"
            results.append(
                KarImportRowResult(row_number=row_number, kar_nummer=kar_nummer, outcome="error", detail=detail)
            )

    db.commit()
    return results


def export_karren_to_xlsx(db: Session) -> bytes:
    """The full KarTracker dataset as a two-sheet workbook: every kar
    (with its status/team/transport-type resolved to names, matching the
    import template's own columns) plus every kar status.
    """
    kar_statuses = {status.id: status.name for status in db.scalars(select(KarTrackerKarStatus)).all()}
    products = {product.id: product.name for product in db.scalars(select(Product)).all()}
    teams = {team.id: team.name for team in db.scalars(select(Team)).all()}

    workbook = Workbook()

    karren_sheet = workbook.active
    karren_sheet.title = "Karren"
    karren_sheet.append(TEMPLATE_COLUMNS)
    for cell in karren_sheet[1]:
        cell.font = Font(bold=True)

    karren = db.scalars(select(KarTrackerKar).order_by(KarTrackerKar.kar_nummer)).all()
    for kar in karren:
        karren_sheet.append(
            [
                kar.kar_nummer,
                kar_statuses.get(kar.status_id, ""),
                teams.get(kar.team_id, "") if kar.team_id is not None else "",
                products.get(kar.transport_type_id, ""),
                kar.last_latitude,
                kar.last_longitude,
                kar.last_recorded_at.replace(tzinfo=None) if kar.last_recorded_at else None,
            ]
        )
    for index, column in enumerate(TEMPLATE_COLUMNS, start=1):
        karren_sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 2, 18)

    statuses_sheet = workbook.create_sheet("KarStatussen")
    statuses_sheet.append(["id", "name"])
    for cell in statuses_sheet[1]:
        cell.font = Font(bold=True)
    for status_id, status_name in sorted(kar_statuses.items(), key=lambda item: item[1]):
        statuses_sheet.append([status_id, status_name])
    statuses_sheet.column_dimensions["A"].width = 8
    statuses_sheet.column_dimensions["B"].width = 24

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
