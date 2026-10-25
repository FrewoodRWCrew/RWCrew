# Atelier's reservations: per festival and per team the persons and the
# package (articles × pieces per person) on one screen; copy from the same
# festival last year; the preparation (persons × package, the same every
# festival day, sizes estimated with the article's split); the picklist PDF;
# the rights, the KPI line and the XLSX tiles; and the article delete rule.

import io
from datetime import date
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models.festival import Festival
from app.db.models.team import Team
from app.modules.module_5.preparation_service import split_by_shares

# The block-1 fixtures and helpers (module, seasons, teams, sizes, a user with every right).
from tests.modules.module_5.test_masterdata import (  # noqa: F401  (fixtures)
    BASE,
    _create_article,
    _login,
    _make_role,
    _make_user,
    _size_id,
    as_atelier,
    world,
)


@pytest.fixture()
def festivals(db_session: Session, world: dict) -> dict:
    """Rock Werchter 2026 (4 days), TW Classic 2026 (1 day) and Rock Werchter 2025."""
    rw = Festival(name="Rock Werchter 2026", start_date=date(2026, 7, 2), end_date=date(2026, 7, 5), season_id=world["season"].id)
    tw = Festival(name="TW Classic", start_date=date(2026, 6, 20), end_date=date(2026, 6, 20), season_id=world["season"].id)
    rw_last = Festival(
        name="Rock Werchter 2025", start_date=date(2025, 7, 3), end_date=date(2025, 7, 6), season_id=world["previous"].id
    )
    db_session.add_all([rw, tw, rw_last])
    db_session.commit()
    return {"rw": rw, "tw": tw, "rw_last": rw_last}


@pytest.fixture()
def articles(as_atelier: TestClient, world: dict) -> dict:
    """A T-shirt in M-women / M-men / L-men (split 20/50/30) and a OneSize apron."""
    m_women, m_men, l_men = _size_id(world, "M", "women"), _size_id(world, "M", "men"), _size_id(world, "L", "men")
    tshirt = as_atelier.post(
        f"{BASE}/articles",
        json={
            "code": "TS",
            "name": "T-shirt",
            "sizes": [
                {"size_id": m_men, "size_share_pct": 50},
                {"size_id": l_men, "size_share_pct": 30},
                {"size_id": m_women, "size_share_pct": 20},
            ],
        },
    ).json()
    apron = _create_article(as_atelier, "SCH", name="Schort", size_ids=[_size_id(world, "OneSize", "unisex")])
    return {"tshirt": tshirt["id"], "apron": apron["id"]}


def _save(client: TestClient, festival_id: int, items: list[dict]):
    """PUT the given team items ({team_id, persons, note?, lines?: [(article, qty)]})."""
    payload = [
        {
            "team_id": item["team_id"],
            "persons": item.get("persons", 0),
            "note": item.get("note"),
            **(
                {"lines": [{"article_id": a, "quantity_per_person": q} for a, q in item["lines"]]}
                if "lines" in item
                else {}
            ),
        }
        for item in items
    ]
    return client.put(f"{BASE}/reservations/{festival_id}", json={"reservations": payload})


def _row(rows: list[dict], team_name: str) -> dict:
    return next(row for row in rows if row["team_name"] == team_name)


def test_split_by_shares_always_adds_up() -> None:
    assert split_by_shares(10, [Decimal(50), Decimal(30), Decimal(20)]) == [5, 3, 2]
    # 7 x (1/3 each): 2.33 -> 2, 2, 2 + the 1 left over to the first.
    assert split_by_shares(7, [Decimal(1)] * 3) == [3, 2, 2]
    assert sum(split_by_shares(13, [Decimal("12.5"), Decimal("37.5"), Decimal(50)])) == 13
    assert split_by_shares(0, [Decimal(1)]) == [0]


def test_size_split_must_add_up_to_100(as_atelier: TestClient, world: dict) -> None:
    m_men, l_men = _size_id(world, "M", "men"), _size_id(world, "L", "men")
    bad = {"code": "X", "name": "X", "sizes": [{"size_id": m_men, "size_share_pct": 50}, {"size_id": l_men}]}
    assert as_atelier.post(f"{BASE}/articles", json=bad).status_code == 400
    good = {"code": "X", "name": "X", "sizes": [{"size_id": m_men, "size_share_pct": 60}, {"size_id": l_men, "size_share_pct": 40}]}
    article = as_atelier.post(f"{BASE}/articles", json=good).json()
    assert [row["size_share_pct"] for row in article["sizes"]] == [60, 40]


def test_every_active_team_is_listed(
    as_atelier: TestClient, db_session: Session, world: dict, festivals: dict
) -> None:
    # Every active team, with or without festival picks or a package yet.
    rows = as_atelier.get(f"{BASE}/reservations", params={"festival_id": festivals["tw"].id}).json()
    assert [(row["team_name"], row["persons"], row["lines"]) for row in rows] == [("RW terras", 0, []), ("VIP", 0, [])]
    # "Oude ploeg" is inactive: not listed, and can't get a reservation.
    old_team = db_session.scalar(select(Team).where(Team.name == "Oude ploeg"))
    assert _save(as_atelier, festivals["tw"].id, [{"team_id": old_team.id, "persons": 2}]).status_code == 400


def test_persons_and_package_saved_together(as_atelier: TestClient, world: dict, festivals: dict, articles: dict) -> None:
    rw, terras = festivals["rw"].id, world["terras"].id
    rows = _save(
        as_atelier, rw, [{"team_id": terras, "persons": 10, "lines": [(articles["tshirt"], 2), (articles["apron"], 1)]}]
    ).json()
    terras_row = _row(rows, "RW terras")
    # Lines come back in article-name order, always 1 per person (whatever was
    # sent: one piece per article at the desk); the changer is the session user.
    assert terras_row["persons"] == 10
    assert terras_row["lines"] == [
        {"article_id": articles["apron"], "quantity_per_person": 1},
        {"article_id": articles["tshirt"], "quantity_per_person": 1},
    ]
    assert terras_row["updated_by_name"] == "atelier@example.com"

    # Without "lines" the package is kept; with lines it's replaced as a whole.
    terras_row = _row(_save(as_atelier, rw, [{"team_id": terras, "persons": 12}]).json(), "RW terras")
    assert (terras_row["persons"], len(terras_row["lines"])) == (12, 2)
    terras_row = _row(_save(as_atelier, rw, [{"team_id": terras, "persons": 12, "lines": [(articles["tshirt"], 3)]}]).json(), "RW terras")
    assert terras_row["lines"] == [{"article_id": articles["tshirt"], "quantity_per_person": 1}]

    # The package belongs to this festival only.
    tw_rows = as_atelier.get(f"{BASE}/reservations", params={"festival_id": festivals["tw"].id}).json()
    assert _row(tw_rows, "RW terras")["lines"] == []

    # Rules: quantity ≥ 1, an article once, no inactive article added anew — and a refused save changes nothing.
    assert _save(as_atelier, rw, [{"team_id": terras, "persons": 1, "lines": [(articles["tshirt"], 0)]}]).status_code == 422
    assert _save(as_atelier, rw, [{"team_id": terras, "persons": 1, "lines": [(articles["tshirt"], 1), (articles["tshirt"], 2)]}]).status_code == 400
    as_atelier.put(f"{BASE}/articles/{articles['apron']}", json={"code": "SCH", "name": "Schort", "active": False})
    response = _save(as_atelier, rw, [{"team_id": terras, "persons": 1, "lines": [(articles["tshirt"], 3), (articles["apron"], 1)]}])
    assert response.status_code == 400
    terras_row = _row(as_atelier.get(f"{BASE}/reservations", params={"festival_id": rw}).json(), "RW terras")
    assert terras_row["persons"] == 12

    # An article used in a package can't be deleted.
    assert as_atelier.delete(f"{BASE}/articles/{articles['tshirt']}").status_code == 400


def test_a_reservation_with_a_package_and_0_persons_stays(
    as_atelier: TestClient, world: dict, festivals: dict, articles: dict
) -> None:
    rw, vip = festivals["rw"].id, world["vip"].id
    _save(as_atelier, rw, [{"team_id": vip, "persons": 0, "lines": [(articles["tshirt"], 1)]}])
    assert _row(as_atelier.get(f"{BASE}/reservations", params={"festival_id": rw}).json(), "VIP")["lines"]
    # Emptying everything removes it again (no "last changed" left).
    _save(as_atelier, rw, [{"team_id": vip, "persons": 0, "lines": []}])
    assert _row(as_atelier.get(f"{BASE}/reservations", params={"festival_id": rw}).json(), "VIP")["updated_at"] is None


def test_copy_from_same_festival_last_year(as_atelier: TestClient, world: dict, festivals: dict, articles: dict) -> None:
    terras, vip = world["terras"].id, world["vip"].id
    cap = _create_article(as_atelier, "CAP", name="Pet")["id"]
    # Last year: terras 8 persons (2 T-shirts + apron), VIP 3 persons (cap).
    _save(as_atelier, festivals["rw_last"].id, [
        {"team_id": terras, "persons": 8, "lines": [(articles["tshirt"], 2), (articles["apron"], 1)]},
        {"team_id": vip, "persons": 3, "lines": [(cap, 1)]},
    ])
    # This year terras already has 10 persons and 1 T-shirt; the cap is no longer active.
    _save(as_atelier, festivals["rw"].id, [{"team_id": terras, "persons": 10, "lines": [(articles["tshirt"], 1)]}])
    as_atelier.put(f"{BASE}/articles/{cap}", json={"code": "CAP", "name": "Pet", "active": False})

    panel = as_atelier.get(f"{BASE}/reservations/{festivals['rw'].id}/previous").json()
    assert panel["festival"]["name"] == "Rock Werchter 2025"
    assert [(row["team_name"], row["persons"]) for row in panel["rows"]] == [("RW terras", 8), ("VIP", 3)]
    # TW Classic has no festival with that name last year.
    assert as_atelier.get(f"{BASE}/reservations/{festivals['tw'].id}/previous").json() == {"festival": None, "rows": []}

    result = as_atelier.post(
        f"{BASE}/reservations/{festivals['rw'].id}/copy-previous", json={"team_ids": [terras, vip]}
    ).json()
    assert result == {"teams_changed": 2, "skipped_inactive": 1}
    rows = as_atelier.get(f"{BASE}/reservations", params={"festival_id": festivals["rw"].id}).json()
    terras_row, vip_row = _row(rows, "RW terras"), _row(rows, "VIP")
    # Terras keeps its own 10 persons; its package gets last year's articles added (1 per person).
    assert terras_row["persons"] == 10
    assert sorted((line["article_id"], line["quantity_per_person"]) for line in terras_row["lines"]) == sorted(
        [(articles["tshirt"], 1), (articles["apron"], 1)]
    )
    # VIP had 0 persons: it takes over last year's 3; the inactive cap is skipped.
    assert (vip_row["persons"], vip_row["lines"]) == (3, [])


def test_preparation_and_picklist(as_atelier: TestClient, world: dict, festivals: dict, articles: dict) -> None:
    rw = festivals["rw"].id
    _save(as_atelier, rw, [
        {"team_id": world["terras"].id, "persons": 10, "lines": [(articles["tshirt"], 2), (articles["apron"], 1)]},
        {"team_id": world["vip"].id, "persons": 3, "lines": [(articles["tshirt"], 1)]},
    ])

    prep = as_atelier.get(f"{BASE}/preparation", params={"festival_id": rw}).json()
    assert len(prep["festival"]["days"]) == 4
    # 1 piece per article per person: 10 + 3 T-shirts and 10 aprons.
    assert (prep["persons_per_day"], prep["pieces_per_day"]) == (13, 23)
    tshirt = next(article for article in prep["articles"] if article["article_code"] == "TS")
    # 13 T-shirts split 20/50/30 over M-women / M-men / L-men: 2.6 / 6.5 / 3.9 → 3 / 6 / 4.
    assert tshirt["quantity"] == 13 and tshirt["equal_split"] is False
    assert [(size["size_name"], size["fit"], size["quantity"]) for size in tshirt["sizes"]] == [
        ("M", "women", 3), ("M", "men", 6), ("L", "men", 4),
    ]
    assert [(bundle["team_name"], bundle["persons"]) for bundle in prep["bundles"]] == [("RW terras", 10), ("VIP", 3)]

    pdf = as_atelier.get(f"{BASE}/preparation/pdf", params={"festival_id": rw, "day": "2026-07-03"})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    one_team = as_atelier.get(
        f"{BASE}/preparation/pdf", params={"festival_id": rw, "day": "2026-07-03", "team_ids": [world["vip"].id]}
    )
    assert one_team.status_code == 200
    assert as_atelier.get(f"{BASE}/preparation/pdf", params={"festival_id": rw, "day": "2026-07-09"}).status_code == 400


def test_kpi_and_rights(
    client: TestClient, as_atelier: TestClient, db_session: Session, world: dict, festivals: dict, articles: dict
) -> None:
    _save(as_atelier, festivals["rw"].id, [
        {"team_id": world["terras"].id, "persons": 10, "lines": [(articles["tshirt"], 1)]},
        {"team_id": world["vip"].id, "persons": 2},
    ])
    stats = as_atelier.get(f"{BASE}/reservation-stats", params={"season_id": world["season"].id}).json()
    assert [(row["festival"]["name"], row["persons_per_day"], row["teams_reserved"], row["teams_without_package"]) for row in stats] == [
        ("TW Classic", 0, 0, 0),
        ("Rock Werchter 2026", 10, 2, 1),
    ]

    # Someone who may only view Voorbereiden can't open or save the reservations.
    as_atelier.post("/api/auth/logout")
    role = _make_role(db_session, "Inpakker", {"atelier.preparation": (True, False, False, False)})
    _make_user(db_session, world, "inpakker@example.com", role)
    _login(client, "inpakker@example.com")
    assert client.get(f"{BASE}/preparation", params={"festival_id": festivals["rw"].id}).status_code == 200
    assert client.get(f"{BASE}/reservations", params={"festival_id": festivals["rw"].id}).status_code == 403
    assert _save(client, festivals["rw"].id, [{"team_id": world["terras"].id, "persons": 1}]).status_code == 403
    assert client.get(f"{BASE}/articles").status_code == 403


def _xlsx(header: list[str], rows: list[list]) -> bytes:
    workbook = Workbook()
    workbook.active.append(header)
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _upload(client: TestClient, path: str, content: bytes) -> list[dict]:
    return client.post(f"{BASE}/{path}-import", files={"file": ("data.xlsx", content)}).json()["results"]


def test_reservation_and_package_line_xlsx(as_atelier: TestClient, world: dict, festivals: dict, articles: dict) -> None:
    results = _upload(as_atelier, "reservations", _xlsx(
        ["season", "festival", "team", "persons"],
        [["2026", "rock werchter 2026", "RW terras", 8], ["2026", "Rock Werchter 2026", "Onbekend", 2]],
    ))
    assert [row["outcome"] for row in results] == ["created", "error"]
    results = _upload(as_atelier, "package-lines", _xlsx(
        ["season", "festival", "team", "article_code", "quantity_per_person"],
        [["2026", "Rock Werchter 2026", "RW terras", "TS", 2], ["2026", "TW Classic", "VIP", "SCH", 1], ["2026", "TW Classic", "VIP", "TS", 0]],
    ))
    # A package line for a team without a reservation yet creates one (0 persons).
    assert [row["outcome"] for row in results] == ["created", "created", "error"]

    terras_row = _row(as_atelier.get(f"{BASE}/reservations", params={"festival_id": festivals["rw"].id}).json(), "RW terras")
    # The uploaded quantity 2 is stored as 1 (one piece per article at the desk).
    assert (terras_row["persons"], terras_row["lines"]) == (8, [{"article_id": articles["tshirt"], "quantity_per_person": 1}])

    export = as_atelier.get(f"{BASE}/package-lines/export")
    rows = [list(row) for row in load_workbook(io.BytesIO(export.content)).active.iter_rows(values_only=True)]
    assert rows[0] == ["season", "festival", "team", "article_code", "quantity_per_person"]
    assert len(rows) == 3
    assert [row["outcome"] for row in _upload(as_atelier, "package-lines", export.content)] == ["unchanged", "unchanged"]


def test_kpi_shows_teams_forgotten_at_another_festival(
    as_atelier: TestClient, db_session: Session, world: dict, festivals: dict
) -> None:
    season = world["season"].id

    def missing() -> dict[str, list[str]]:
        stats = as_atelier.get(f"{BASE}/reservation-stats", params={"season_id": season}).json()
        return {row["festival"]["name"]: row["teams_missing"] for row in stats}

    # Nobody reserved anywhere: nobody is "forgotten".
    assert missing() == {"TW Classic": [], "Rock Werchter 2026": []}

    # Terras reserved at Rock Werchter only → probably forgotten at TW Classic (not at RW, not last season).
    _save(as_atelier, festivals["rw"].id, [{"team_id": world["terras"].id, "persons": 6}])
    assert missing() == {"TW Classic": ["RW terras"], "Rock Werchter 2026": []}

    # 0 persons with a note ("komt niet") silences it.
    _save(as_atelier, festivals["tw"].id, [{"team_id": world["terras"].id, "persons": 0, "note": "komt niet"}])
    assert missing() == {"TW Classic": [], "Rock Werchter 2026": []}

    # An inactive team is never reported.
    _save(as_atelier, festivals["rw"].id, [{"team_id": world["vip"].id, "persons": 2}])
    assert missing()["TW Classic"] == ["VIP"]
    world["vip"].active = False
    db_session.commit()
    assert missing()["TW Classic"] == []
