# These tests check the product half of StockMaster's "Data Upload/Download"
# screen (moved here from MasterData together with the product screens): a
# single "stockmaster.dataupload" screen key gates the XLSX template
# download, bulk import and export of products and their four lookup lists.
# Product is the interesting shape (multiple FK lookups + booleans, no
# duplicate rule); the lookups share MasterData's simple name-only shape.

import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.db.models.module import Module
from app.db.models.product import Product
from app.db.models.product_category import ProductCategory
from app.db.models.product_limit import ProductLimit
from app.db.models.product_type import ProductType
from app.db.models.stockmaster_role import StockMasterRole
from app.db.models.stockmaster_role_permission import StockMasterRolePermission
from app.db.models.stockmaster_screen import StockMasterScreen
from app.db.models.stockmaster_user_role import StockMasterUserRole
from app.db.models.user import User
from app.db.models.user_module_access import UserModuleAccess
from app.db.models.warehouse import Warehouse
from app.modules.module_4.screens import sync_screens

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _build_xlsx(header: list[str], rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _create_user(db_session: Session, *, email: str, is_super_admin: bool = False) -> User:
    user = User(
        email=email,
        hashed_password=hash_password("password123"),
        display_name=email,
        is_super_admin=is_super_admin,
    )
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return user


def _create_stockmaster_module(db_session: Session) -> Module:
    module = db_session.scalar(select(Module).where(Module.key == "module-4"))
    if module is not None:
        return module
    module = Module(key="module-4", name="StockMaster", sort_order=4)
    db_session.add(module)
    db_session.commit()
    db_session.refresh(module)
    return module


def _grant_module_access(db_session: Session, user: User, module: Module) -> None:
    db_session.add(UserModuleAccess(user_id=user.id, module_id=module.id))
    db_session.commit()


def _grant_screen_permission(
    db_session: Session,
    user: User,
    screen_key: str,
    *,
    can_view: bool = False,
    can_create: bool = False,
) -> None:
    role = StockMasterRole(name=f"{screen_key} role for {user.email}")
    db_session.add(role)
    db_session.commit()
    db_session.refresh(role)

    screen = db_session.scalar(select(StockMasterScreen).where(StockMasterScreen.key == screen_key))
    db_session.add(
        StockMasterRolePermission(
            role_id=role.id,
            screen_id=screen.id,
            can_view=can_view,
            can_create=can_create,
            can_edit=False,
            can_delete=False,
        )
    )
    db_session.add(StockMasterUserRole(user_id=user.id, role_id=role.id))
    db_session.commit()


def _login(client: TestClient, email: str) -> None:
    client.post("/api/auth/login", json={"email": email, "password": "password123"})


def _login_as_admin(client: TestClient, db_session: Session) -> Module:
    sync_screens(db_session)
    module = _create_stockmaster_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, module)
    _login(client, "admin@example.com")
    return module


# The lookup columns every imported product needs (type, warehouse, category, Limiet).
_PRODUCT_LOOKUP_HEADER = ["type_name", "warehouse_name", "category_name", "limit_name"]
_PRODUCT_LOOKUP_CELLS = ["Drinks", "Main", "Bar", "Ongelimiteerd"]


def _create_product_lookups(db_session: Session) -> None:
    """The lookup rows _PRODUCT_LOOKUP_CELLS refer to, plus a "Percentage" Limiet."""
    db_session.add_all(
        [
            ProductType(name="Drinks"),
            Warehouse(name="Main"),
            ProductCategory(name="Bar"),
            ProductLimit(name="Ongelimiteerd"),
            ProductLimit(name="Percentage"),
        ]
    )
    db_session.commit()


def test_product_import_resolves_lookups_and_booleans(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    _create_product_lookups(db_session)

    xlsx_bytes = _build_xlsx(
        [
            "name",
            "type_name",
            "warehouse_name",
            "warehouse_location",
            "category_name",
            "is_consumable",
            "is_blocked",
            "is_logistics_product",
            "limit_name",
            "description",
        ],
        [
            ["Water bottles", "Drinks", "Main", "Shelf A1", "Bar", "true", "", "", "Ongelimiteerd", "Case of 24"],
            ["Bad product", "Nonexistent Type", "Main", "", "Bar", "", "", "", "Ongelimiteerd", ""],
        ],
    )
    response = client.post(
        "/api/modules/module-4/product-import", files={"file": ("products.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)}
    )

    assert response.status_code == 200
    results = {row["name"]: row for row in response.json()["results"]}
    assert results["Water bottles"]["outcome"] == "created"
    assert results["Bad product"]["outcome"] == "error"

    product = db_session.scalar(select(Product).where(Product.name == "Water bottles"))
    assert product is not None
    assert product.is_consumable is True
    assert product.is_blocked is False
    assert product.warehouse_location == "Shelf A1"


def test_product_import_rejects_missing_lookups_percentage_without_value_and_huge_cost(
    client: TestClient, db_session: Session
) -> None:
    _login_as_admin(client, db_session)
    _create_product_lookups(db_session)

    xlsx_bytes = _build_xlsx(
        ["name", *_PRODUCT_LOOKUP_HEADER, "limit_value", "cost"],
        [
            ["No type", "", "Main", "Bar", "Ongelimiteerd", "", ""],
            ["No limit", "Drinks", "Main", "Bar", "", "", ""],
            ["Percentage without value", "Drinks", "Main", "Bar", "Percentage", "", ""],
            ["Huge cost", "Drinks", "Main", "Bar", "Ongelimiteerd", "", "100000000"],
            ["Percentage with value", "Drinks", "Main", "Bar", "Percentage", "50", "12.346"],
        ],
    )
    response = client.post(
        "/api/modules/module-4/product-import", files={"file": ("products.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)}
    )

    assert response.status_code == 200
    results = {row["name"]: row for row in response.json()["results"]}
    assert "type_name" in results["No type"]["detail"]
    assert "limit_name" in results["No limit"]["detail"]
    assert "limit_value" in results["Percentage without value"]["detail"]
    assert "cost" in results["Huge cost"]["detail"]
    assert results["Percentage with value"]["outcome"] == "created"
    product = db_session.scalar(select(Product).where(Product.name == "Percentage with value"))
    assert product.limit_value == 50
    assert str(product.cost) == "12.35"


def test_product_import_rejects_an_unrecognized_boolean_value(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    _create_product_lookups(db_session)

    xlsx_bytes = _build_xlsx(
        ["name", *_PRODUCT_LOOKUP_HEADER, "is_consumable"],
        [["Typo Product", *_PRODUCT_LOOKUP_CELLS, "tru"]],
    )
    response = client.post(
        "/api/modules/module-4/product-import", files={"file": ("products.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)}
    )

    assert response.status_code == 200
    result = response.json()["results"][0]
    assert result["outcome"] == "error"
    assert "is_consumable" in result["detail"]
    assert db_session.scalar(select(Product).where(Product.name == "Typo Product")) is None


def test_product_import_allows_duplicate_names(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    _create_product_lookups(db_session)
    db_session.add(Product(name="Water bottles"))
    db_session.commit()

    xlsx_bytes = _build_xlsx(["name", *_PRODUCT_LOOKUP_HEADER], [["Water bottles", *_PRODUCT_LOOKUP_CELLS]])
    response = client.post(
        "/api/modules/module-4/product-import", files={"file": ("products.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)}
    )

    assert response.status_code == 200
    assert response.json()["results"][0]["outcome"] == "created"
    assert len(db_session.scalars(select(Product).where(Product.name == "Water bottles")).all()) == 2


# --- Permissions --------------------------------------------------------------


def test_product_template_needs_dataupload_view_not_the_products_screen(
    client: TestClient, db_session: Session
) -> None:
    sync_screens(db_session)
    module = _create_stockmaster_module(db_session)
    user = _create_user(db_session, email="member@example.com")
    _grant_module_access(db_session, user, module)
    _grant_screen_permission(db_session, user, "stockmaster.products", can_view=True, can_create=True)
    _login(client, "member@example.com")

    assert client.get("/api/modules/module-4/product-import/template").status_code == 403


def test_view_only_dataupload_permission_cannot_import_warehouses(client: TestClient, db_session: Session) -> None:
    sync_screens(db_session)
    module = _create_stockmaster_module(db_session)
    user = _create_user(db_session, email="viewer@example.com")
    _grant_module_access(db_session, user, module)
    _grant_screen_permission(db_session, user, "stockmaster.dataupload", can_view=True)
    _login(client, "viewer@example.com")

    assert client.get("/api/modules/module-4/warehouse-import/template").status_code == 200
    xlsx_bytes = _build_xlsx(["id", "name"], [[None, "Hal B"]])
    response = client.post(
        "/api/modules/module-4/warehouse-import", files={"file": ("w.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)}
    )
    assert response.status_code == 403


def test_product_export_contains_every_product(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    db_session.add_all([Product(name="Tape"), Product(name="Lader")])
    db_session.commit()

    response = client.get("/api/modules/module-4/products/export")

    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    names = {row[1] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert {"Tape", "Lader"} <= names


def _upload(client: TestClient, path: str, header: list[str], rows: list[list]) -> list[dict]:
    """Upload a small workbook and return its per-row results."""
    response = client.post(
        f"/api/modules/module-4/{path}", files={"file": ("upload.xlsx", _build_xlsx(header, rows), XLSX_CONTENT_TYPE)}
    )
    assert response.status_code == 200, response.text
    return response.json()["results"]


# --- Update through the "id" column -------------------------------------------


def test_lookup_row_with_an_id_renames_and_unknown_id_is_an_error(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    main, spare = Warehouse(name="Main"), Warehouse(name="Spare")
    db_session.add_all([main, spare])
    db_session.commit()

    results = _upload(
        client,
        "warehouse-import",
        ["id", "name"],
        [[main.id, "Central"], [None, "North"], [spare.id, "Central"], [999, "Ghost"], ["abc", "Bad"]],
    )

    assert [(row["name"], row["outcome"]) for row in results] == [
        ("Central", "updated"),
        ("North", "created"),
        ("Central", "error"),
        ("Ghost", "error"),
        ("Bad", "error"),
    ]
    assert results[2]["detail"] == "A warehouse with this name already exists"
    assert results[3]["detail"] == "id 999 not found"
    db_session.refresh(main)
    assert main.name == "Central"


def test_product_export_reuploaded_updates_instead_of_duplicating(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    _create_product_lookups(db_session)
    created = _upload(
        client,
        "product-import",
        ["name", "type_name", "warehouse_name", "category_name", "limit_name", "min_stock", "stock_return_to"],
        [["Cola", "Drinks", "Main", "Bar", "Ongelimiteerd", 12, "Free"]],
    )
    assert [row["outcome"] for row in created] == ["created"]
    product = db_session.scalar(select(Product).where(Product.name == "Cola"))
    assert (product.min_stock, product.stock_return_to) == (12, "free")

    # The export has both new columns; uploaded again as-is, it updates.
    exported = client.get(f"/api/modules/module-4/products/export").content
    header = [cell.value for cell in load_workbook(io.BytesIO(exported)).active[1]]
    assert {"min_stock", "stock_return_to"} <= set(header)
    again = client.post(f"/api/modules/module-4/product-import", files={"file": ("p.xlsx", exported, XLSX_CONTENT_TYPE)})
    assert [row["outcome"] for row in again.json()["results"]] == ["updated"]
    assert len(db_session.scalars(select(Product).where(Product.name == "Cola")).all()) == 1

    # Only the columns in the sheet change; a bad stock_return_to is a row error.
    results = _upload(
        client, "product-import", ["id", "min_stock", "stock_return_to"], [[product.id, 3, None], [product.id, 1, "x"]]
    )
    assert [row["outcome"] for row in results] == ["updated", "error"]
    db_session.refresh(product)
    assert (product.name, product.min_stock, product.stock_return_to) == ("Cola", 3, "kar")
