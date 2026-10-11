# Inboeken from Excel (app/modules/module_4/booking_import.py): the template
# lists the bookable products, and an uploaded copy becomes basket lines
# plus per-row problems — without booking anything.

import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.db.models.product import Product
from tests.modules.module_4.conftest import BASE, StockWorld, login, make_role, stock_of

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _upload(client: TestClient, rows: list[list[object]]):
    """POST a workbook with the template header and the given rows."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["product_id", "product_name", "bin", "quantity"])
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return client.post(f"{BASE}/book-in-import", files={"file": ("inboeken.xlsx", buffer.getvalue(), XLSX)})


def test_template_lists_every_bookable_product(as_warehouse: TestClient, world: StockWorld) -> None:
    response = as_warehouse.get(f"{BASE}/book-in-import/template")

    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("product_id", "product_name", "bin", "quantity")
    # Sorted by name, blocked products ("Oude kabel") left out, quantity empty.
    assert rows[1:] == [
        (world.tape.id, "Tape", "Hal A · A-01", None),
        (world.lader.id, "Walkie-lader", "Hal A · B-12", None),
    ]


def test_upload_returns_basket_lines_and_row_problems(as_warehouse: TestClient, world: StockWorld) -> None:
    response = _upload(
        as_warehouse,
        [
            [world.lader.id, "Walkie-lader", None, 5],  # row 2: by id
            [None, "tape", None, 3],                     # row 3: by name, any case
            [world.tape.id, "Tape", None, None],         # row 4: no quantity → skipped
            [world.lader.id, "Walkie-lader", None, 2],   # row 5: same product again → summed
            [None, "Bestaat niet", None, 1],             # row 6: unknown
            [world.tape.id, "Tape", None, "veel"],       # row 7: not a number
            [world.tape.id, "Tape", None, 0],            # row 8: zero → skipped
        ],
    )

    assert response.status_code == 200
    body = response.json()
    assert body["lines"] == [
        {"product_id": world.lader.id, "quantity": 7},
        {"product_id": world.tape.id, "quantity": 3},
    ]
    assert body["errors"] == [
        {"row_number": 6, "detail": "unknown_product|Bestaat niet"},
        {"row_number": 7, "detail": "bad_quantity|veel"},
    ]
    # Nothing was booked.
    assert stock_of(as_warehouse, world.lader.id)["total"] == 0


def test_ambiguous_name_is_a_row_problem(as_warehouse: TestClient, world: StockWorld, db_session: Session) -> None:
    db_session.add(Product(name="Tape"))
    db_session.commit()

    body = _upload(as_warehouse, [[None, "Tape", None, 1]]).json()

    assert body == {"lines": [], "errors": [{"row_number": 2, "detail": "ambiguous_product|Tape"}]}


def test_unreadable_file_is_refused(as_warehouse: TestClient) -> None:
    response = as_warehouse.post(f"{BASE}/book-in-import", files={"file": ("x.xlsx", b"not an excel", XLSX)})

    assert response.status_code == 400
    assert response.json()["detail"] == "import_unreadable"


def test_import_needs_the_book_in_right(client: TestClient, world: StockWorld, make_user, db_session: Session) -> None:
    viewer = make_role(db_session, "Kijker", {"stockmaster.stock": (True, False, False, False)})
    make_user("kijker@example.com", viewer)
    login(client, "kijker@example.com")

    assert client.get(f"{BASE}/book-in-import/template").status_code == 403
    assert _upload(client, [[world.lader.id, "Walkie-lader", None, 1]]).status_code == 403
