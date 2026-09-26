# These tests check MasterData's "Data Upload/Download" screen: a single
# "masterdata.dataupload" screen key gates the XLSX template download,
# bulk import, and export endpoints for every one of the twelve
# master-data tables (mirrors KarTracker's own
# tests/modules/module_2/test_kar_import.py /
# test_kar_status_import.py). Coverage here focuses on one simple lookup
# (Season), Team (unique name + optional FK lookups, scalar fields only),
# and Product (multiple optional FK lookups + booleans, no duplicate
# rule) — the three distinct shapes every other table's import file
# follows — plus permission-independence checks that apply to all of them.
#
# Local helper functions are used throughout (not a shared conftest),
# matching every other test file already in this directory — module_9,
# unlike module_2, has no tests/modules/module_9/conftest.py.

import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.db.models.masterdata_role import MasterDataRole
from app.db.models.masterdata_role_permission import MasterDataRolePermission
from app.db.models.masterdata_screen import MasterDataScreen
from app.db.models.masterdata_user_role import MasterDataUserRole
from app.db.models.module import Module
from app.db.models.product import Product
from app.db.models.season import Season
from app.db.models.team import Team
from app.db.models.team_location import TeamLocation
from app.db.models.user import User
from app.db.models.user_module_access import UserModuleAccess
from app.modules.module_9.screens import sync_screens

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _build_xlsx(header: list[str], rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _create_user(db_session: Session, *, email: str, is_super_admin: bool = False) -> User:
    user = User(
        email=email,
        hashed_password=hash_password("password123"),
        display_name=email,
        is_super_admin=is_super_admin,
    )
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return user


def _create_masterdata_module(db_session: Session) -> Module:
    module = db_session.scalar(select(Module).where(Module.key == "module-9"))
    if module is not None:
        return module
    module = Module(key="module-9", name="MasterData", sort_order=9)
    db_session.add(module)
    db_session.commit()
    db_session.refresh(module)
    return module


def _grant_module_access(db_session: Session, user: User, module: Module) -> None:
    db_session.add(UserModuleAccess(user_id=user.id, module_id=module.id))
    db_session.commit()


def _grant_screen_permission(
    db_session: Session,
    user: User,
    screen_key: str,
    *,
    can_view: bool = False,
    can_create: bool = False,
) -> None:
    role = MasterDataRole(name=f"{screen_key} role for {user.email}")
    db_session.add(role)
    db_session.commit()
    db_session.refresh(role)

    screen = db_session.scalar(select(MasterDataScreen).where(MasterDataScreen.key == screen_key))
    db_session.add(
        MasterDataRolePermission(
            role_id=role.id,
            screen_id=screen.id,
            can_view=can_view,
            can_create=can_create,
            can_edit=False,
            can_delete=False,
        )
    )
    db_session.add(MasterDataUserRole(user_id=user.id, role_id=role.id))
    db_session.commit()


def _login(client: TestClient, email: str) -> None:
    client.post("/api/auth/login", json={"email": email, "password": "password123"})


def _login_as_admin(client: TestClient, db_session: Session) -> Module:
    sync_screens(db_session)
    module = _create_masterdata_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, module)
    _login(client, "admin@example.com")
    return module


# --- Season: the simple name-only lookup shape ------------------------------


def test_downloading_the_season_template_requires_dataupload_view_permission(
    client: TestClient, db_session: Session
) -> None:
    sync_screens(db_session)
    module = _create_masterdata_module(db_session)
    user = _create_user(db_session, email="member@example.com")
    _grant_module_access(db_session, user, module)
    _grant_screen_permission(db_session, user, "masterdata.season", can_view=True)
    _login(client, "member@example.com")

    response = client.get("/api/modules/module-9/season-import/template")

    assert response.status_code == 403


def test_super_admin_can_download_the_season_template(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)

    response = client.get("/api/modules/module-9/season-import/template")

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_CONTENT_TYPE
    assert "attachment" in response.headers["content-disposition"]

    workbook = load_workbook(io.BytesIO(response.content))
    header_row = [cell.value for cell in workbook.active[1]]
    assert header_row == ["name"]


def test_view_only_dataupload_permission_cannot_import_seasons(client: TestClient, db_session: Session) -> None:
    sync_screens(db_session)
    module = _create_masterdata_module(db_session)
    user = _create_user(db_session, email="viewer@example.com")
    _grant_module_access(db_session, user, module)
    _grant_screen_permission(db_session, user, "masterdata.dataupload", can_view=True)
    _login(client, "viewer@example.com")

    xlsx_bytes = _build_xlsx(["name"], [["2027"]])
    response = client.post(
        "/api/modules/module-9/season-import", files={"file": ("seasons.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)}
    )

    assert response.status_code == 403


def test_season_import_creates_new_rows_and_rejects_duplicates(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    existing = Season(name="2026")
    db_session.add(existing)
    db_session.commit()

    xlsx_bytes = _build_xlsx(["name"], [["2027"], ["2026"]])
    response = client.post(
        "/api/modules/module-9/season-import", files={"file": ("seasons.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)}
    )

    assert response.status_code == 200
    outcomes = {row["name"]: row["outcome"] for row in response.json()["results"]}
    assert outcomes["2027"] == "created"
    assert outcomes["2026"] == "error"

    assert db_session.scalar(select(Season).where(Season.name == "2027")) is not None
    # No second "2026" row was created for the duplicate.
    seasons_named_2026 = db_session.scalars(select(Season).where(Season.name == "2026")).all()
    assert len(seasons_named_2026) == 1


def test_season_export_contains_every_season(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    season = Season(name="2026")
    db_session.add(season)
    db_session.commit()
    db_session.refresh(season)

    response = client.get("/api/modules/module-9/seasons/export")

    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    rows = list(workbook.active.iter_rows(values_only=True))
    assert rows[0] == ("id", "name")
    assert rows[1] == (season.id, season.name)


def test_dataupload_permission_is_independent_of_season_screen_permission(
    client: TestClient, db_session: Session
) -> None:
    sync_screens(db_session)
    module = _create_masterdata_module(db_session)
    user = _create_user(db_session, email="viewer@example.com")
    _grant_module_access(db_session, user, module)
    # Full view+create on Season's own screen, but nothing on dataupload.
    _grant_screen_permission(db_session, user, "masterdata.season", can_view=True, can_create=True)
    _login(client, "viewer@example.com")

    xlsx_bytes = _build_xlsx(["name"], [["2027"]])
    response = client.post(
        "/api/modules/module-9/season-import", files={"file": ("seasons.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)}
    )

    assert response.status_code == 403


# --- Team: unique name + optional FK-by-name lookups, scalar fields only ----


def test_team_import_resolves_location_and_rejects_duplicate_names(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    location = TeamLocation(location="Main Stage")
    db_session.add(location)
    existing_team = Team(name="Logistics")
    db_session.add(existing_team)
    db_session.commit()

    xlsx_bytes = _build_xlsx(
        ["name", "location_name", "delivery_method_name", "description", "active"],
        [
            ["Bar Team", "Main Stage", "", "Serves drinks", "no"],
            ["Logistics", "", "", ""],
            ["Ghost Team", "Nonexistent Location", "", ""],
        ],
    )
    response = client.post(
        "/api/modules/module-9/team-import", files={"file": ("teams.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)}
    )

    assert response.status_code == 200
    results = {row["name"]: row for row in response.json()["results"]}
    assert results["Bar Team"]["outcome"] == "created"
    assert results["Logistics"]["outcome"] == "error"
    assert results["Ghost Team"]["outcome"] == "error"
    assert "Nonexistent Location" in results["Ghost Team"]["detail"]

    new_team = db_session.scalar(select(Team).where(Team.name == "Bar Team"))
    assert new_team is not None
    assert new_team.location_id == location.id
    assert new_team.active is False


def test_team_export_does_not_include_task_or_kernlid_columns(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    team = Team(name="Logistics")
    db_session.add(team)
    db_session.commit()

    response = client.get("/api/modules/module-9/teams/export")

    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    header_row = [cell.value for cell in workbook.active[1]]
    assert header_row == ["id", "name", "location_name", "delivery_method_name", "description", "active"]


# --- Product: several optional FK-by-name lookups + booleans, no dup rule --


def test_product_import_resolves_optional_lookups_and_booleans(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)

    xlsx_bytes = _build_xlsx(
        [
            "name",
            "type_name",
            "warehouse_name",
            "warehouse_location",
            "category_name",
            "is_consumable",
            "is_blocked",
            "is_logistics_product",
            "limit_name",
            "description",
        ],
        [
            ["Water bottles", "", "", "Shelf A1", "", "true", "", "", "", "Case of 24"],
            ["Bad product", "Nonexistent Type", "", "", "", "", "", "", "", ""],
        ],
    )
    response = client.post(
        "/api/modules/module-9/product-import", files={"file": ("products.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)}
    )

    assert response.status_code == 200
    results = {row["name"]: row for row in response.json()["results"]}
    assert results["Water bottles"]["outcome"] == "created"
    assert results["Bad product"]["outcome"] == "error"

    product = db_session.scalar(select(Product).where(Product.name == "Water bottles"))
    assert product is not None
    assert product.is_consumable is True
    assert product.is_blocked is False
    assert product.warehouse_location == "Shelf A1"


def test_product_import_rejects_an_unrecognized_boolean_value(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)

    xlsx_bytes = _build_xlsx(
        ["name", "is_consumable"],
        [["Typo Product", "tru"]],
    )
    response = client.post(
        "/api/modules/module-9/product-import", files={"file": ("products.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)}
    )

    assert response.status_code == 200
    result = response.json()["results"][0]
    assert result["outcome"] == "error"
    assert "is_consumable" in result["detail"]
    assert db_session.scalar(select(Product).where(Product.name == "Typo Product")) is None


def test_product_import_allows_duplicate_names(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    db_session.add(Product(name="Water bottles"))
    db_session.commit()

    xlsx_bytes = _build_xlsx(["name"], [["Water bottles"]])
    response = client.post(
        "/api/modules/module-9/product-import", files={"file": ("products.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)}
    )

    assert response.status_code == 200
    assert response.json()["results"][0]["outcome"] == "created"
    assert len(db_session.scalars(select(Product).where(Product.name == "Water bottles")).all()) == 2


# --- get_my_permissions: creatable_screen_keys -----------------------------


def test_my_permissions_reports_dataupload_as_creatable_when_granted(client: TestClient, db_session: Session) -> None:
    sync_screens(db_session)
    module = _create_masterdata_module(db_session)
    user = _create_user(db_session, email="uploader@example.com")
    _grant_module_access(db_session, user, module)
    _grant_screen_permission(db_session, user, "masterdata.dataupload", can_view=True, can_create=True)
    _login(client, "uploader@example.com")

    response = client.get("/api/modules/module-9/me/permissions")

    assert response.status_code == 200
    body = response.json()
    assert "masterdata.dataupload" in body["viewable_screen_keys"]
    assert "masterdata.dataupload" in body["creatable_screen_keys"]


def test_my_permissions_does_not_report_dataupload_as_creatable_for_view_only_user(
    client: TestClient, db_session: Session
) -> None:
    sync_screens(db_session)
    module = _create_masterdata_module(db_session)
    user = _create_user(db_session, email="viewer@example.com")
    _grant_module_access(db_session, user, module)
    _grant_screen_permission(db_session, user, "masterdata.dataupload", can_view=True)
    _login(client, "viewer@example.com")

    response = client.get("/api/modules/module-9/me/permissions")

    assert response.status_code == 200
    body = response.json()
    assert "masterdata.dataupload" in body["viewable_screen_keys"]
    assert "masterdata.dataupload" not in body["creatable_screen_keys"]
