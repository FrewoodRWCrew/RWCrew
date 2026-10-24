# These tests check the parts of MasterData's "Data Upload/Download" added
# for keeping the uploads in line with the tables: updating a record through
# its "id" column (so an export can be edited and uploaded again), the new
# Season columns, and the four team-link tiles (tasks, Kernleden and
# festivals per team, Ploegverantwoordelijken). Helpers come from
# test_data_upload_download.py, like the scanner tests do in module_1.

import io
from datetime import date

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models.altsien_select_step_progress import AltsienSelectStepProgress
from app.db.models.festival import Festival
from app.db.models.season import Season
from app.db.models.team import Team
from app.db.models.team_festival import TeamFestival
from app.db.models.team_kernlid import TeamKernlid
from app.db.models.team_responsible import TeamResponsible
from app.db.models.team_task import TeamTask
from app.db.models.team_team_task import TeamTeamTask
from tests.modules.module_9.test_data_upload_download import (
    XLSX_CONTENT_TYPE,
    _build_xlsx,
    _create_user,
    _login_as_admin,
)

BASE = "/api/modules/module-9"


def _upload(client: TestClient, path: str, header: list[str], rows: list[list]) -> list[dict]:
    """Upload a small workbook and return its per-row results."""
    response = client.post(
        f"{BASE}/{path}", files={"file": ("upload.xlsx", _build_xlsx(header, rows), XLSX_CONTENT_TYPE)}
    )
    assert response.status_code == 200, response.text
    return response.json()["results"]


def _export_rows(client: TestClient, path: str) -> list[tuple]:
    """Every row (header first) of an export."""
    response = client.get(f"{BASE}/{path}")
    assert response.status_code == 200
    return list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))


# --- Update through the "id" column ---------------------------------------


def test_season_periode_open_round_trips_and_updates(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    season = Season(name="2026", periode_open=False)
    db_session.add(season)
    db_session.commit()

    results = _upload(
        client, "season-import", ["id", "name", "periode_open"], [[season.id, "2026", "yes"], [None, "2027", None]]
    )

    assert [row["outcome"] for row in results] == ["updated", "created"]
    db_session.refresh(season)
    assert season.periode_open is True
    new_season = db_session.scalar(select(Season).where(Season.name == "2027"))
    assert new_season.periode_open is False
    assert _export_rows(client, "seasons/export")[1:] == [(season.id, "2026", "yes"), (new_season.id, "2027", "no")]


# --- The team links ---------------------------------------------------------


def test_team_task_links_are_added_once_and_exported(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    team, task = Team(name="Logistics"), TeamTask(team_tasks="Opbouw")
    db_session.add_all([team, task])
    db_session.commit()

    results = _upload(
        client,
        "team-task-link-import",
        ["team_name", "team_task"],
        [["logistics", "opbouw"], ["Logistics", "Opbouw"], ["Logistics", "Afbraak"]],
    )

    assert [row["outcome"] for row in results] == ["created", "unchanged", "error"]
    assert results[0]["key"] == "Logistics — Opbouw"
    assert results[2]["detail"] == 'Team task "Afbraak" not found'
    assert db_session.scalar(select(TeamTeamTask.team_task_id).where(TeamTeamTask.team_id == team.id)) == task.id
    assert _export_rows(client, "team-task-links/export") == [("team_name", "team_task"), ("Logistics", "Opbouw")]


def test_kernlid_links_need_an_altsien_kernlid_user(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    team = Team(name="Logistics")
    kernlid = _create_user(db_session, email="kern@example.com")
    kernlid.is_altsien_kernlid = True
    _create_user(db_session, email="other@example.com")
    db_session.add(team)
    db_session.commit()

    results = _upload(
        client,
        "team-kernlid-import",
        ["team_name", "user_email"],
        [["Logistics", "KERN@example.com"], ["Logistics", "other@example.com"]],
    )

    assert [row["outcome"] for row in results] == ["created", "error"]
    assert results[1]["detail"] == 'User "other@example.com" is not an Altsien Kernlid'
    assert db_session.scalar(select(TeamKernlid.altsien_kernlid_id)) == kernlid.id


def test_festival_links_follow_the_wizard_rules(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    season, team = Season(name="2026"), Team(name="Logistics")
    db_session.add_all([season, team])
    db_session.commit()
    day = date(2026, 7, 1)
    db_session.add_all(
        [
            Festival(name="Rock Werchter", start_date=day, end_date=day, season_id=season.id, active=True),
            Festival(name="Closed Fest", start_date=day, end_date=day, season_id=season.id, active=False),
            AltsienSelectStepProgress(season_id=season.id, team_id=team.id, step_key="afleverlocaties"),
        ]
    )
    db_session.commit()

    results = _upload(
        client,
        "team-festival-import",
        ["season_name", "team_name", "festival_name"],
        [["2026", "Logistics", "rock werchter"], ["2026", "Logistics", "Closed Fest"]],
    )

    assert [row["outcome"] for row in results] == ["created", "error"]
    assert db_session.scalar(select(TeamFestival.team_id)) == team.id
    # A new festival means the delivery locations must be confirmed again.
    assert db_session.scalar(select(AltsienSelectStepProgress.id)) is None


def test_team_responsibles_are_validated_and_updated_by_id(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    season, team = Season(name="2026"), Team(name="Logistics")
    db_session.add_all([season, team])
    db_session.commit()
    header = ["id", "season_name", "team_name", "name", "email", "phone", "comments"]

    results = _upload(
        client,
        "team-responsible-import",
        header,
        [
            [None, "2026", "Logistics", "An", "an@example.com", "0470", None],
            [None, "2026", "Logistics", "Bo", "nope", "1", None],
        ],
    )
    assert [row["outcome"] for row in results] == ["created", "error"]
    assert results[1]["detail"].startswith("email:")

    person = db_session.scalar(select(TeamResponsible))
    updated = _upload(client, "team-responsible-import", ["id", "phone"], [[person.id, "0499"]])
    assert [row["outcome"] for row in updated] == ["updated"]
    db_session.refresh(person)
    assert (person.name, person.email, person.phone) == ("An", "an@example.com", "0499")
    assert _export_rows(client, "team-responsibles/export")[1] == (
        person.id,
        "2026",
        "Logistics",
        "An",
        "an@example.com",
        "0499",
        None,
    )


def test_link_endpoints_need_a_login(client: TestClient) -> None:
    for path in ("team-task-link-import", "team-kernlid-import", "team-festival-import", "team-responsible-import"):
        assert client.get(f"{BASE}/{path}/template").status_code == 401
