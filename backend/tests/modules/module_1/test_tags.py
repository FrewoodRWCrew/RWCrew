# These tests check the "TagManagement" screen: gated by normal module
# access + the "tagscan.tag-management" permission, independently of
# "tagscan.dashboard"/"roles"/"users", with a super-admin bypass and full
# CRUD — mirrors module_9's test_products.py shape.

import io
from datetime import date, datetime

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.db.models.module import Module
from app.db.models.product import Product
from app.db.models.rfid_tag import RfidTag
from app.db.models.tagscan_role import TagscanRole
from app.db.models.tagscan_role_permission import TagscanRolePermission
from app.db.models.tagscan_screen import TagscanScreen
from app.db.models.tagscan_user_role import TagscanUserRole
from app.db.models.user import User
from app.db.models.user_module_access import UserModuleAccess
from app.modules.module_1.screens import sync_screens as sync_tagscan_screens


def _create_user(db_session: Session, *, email: str, is_super_admin: bool = False) -> User:
    user = User(
        email=email,
        hashed_password=hash_password("password123"),
        display_name=email,
        is_super_admin=is_super_admin,
    )
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return user


def _create_tagscan_module(db_session: Session) -> Module:
    module = db_session.scalar(select(Module).where(Module.key == "module-1"))
    if module is not None:
        return module
    module = Module(key="module-1", name="Tagscan", sort_order=1)
    db_session.add(module)
    db_session.commit()
    db_session.refresh(module)
    return module


def _grant_module_access(db_session: Session, user: User, module: Module) -> None:
    db_session.add(UserModuleAccess(user_id=user.id, module_id=module.id))
    db_session.commit()


def _grant_tag_management_permission(
    db_session: Session,
    user: User,
    *,
    can_view: bool = False,
    can_create: bool = False,
    can_edit: bool = False,
    can_delete: bool = False,
) -> None:
    role = TagscanRole(name=f"TagManagement Role {user.email}")
    db_session.add(role)
    db_session.commit()
    db_session.refresh(role)

    screen = db_session.scalar(select(TagscanScreen).where(TagscanScreen.key == "tagscan.tag-management"))
    db_session.add(
        TagscanRolePermission(
            role_id=role.id,
            screen_id=screen.id,
            can_view=can_view,
            can_create=can_create,
            can_edit=can_edit,
            can_delete=can_delete,
        )
    )
    db_session.add(TagscanUserRole(user_id=user.id, role_id=role.id))
    db_session.commit()


def _create_tag(db_session: Session, *, epc_uid: str = "E200001122334455") -> RfidTag:
    tag = RfidTag(epc_uid=epc_uid)
    db_session.add(tag)
    db_session.commit()
    db_session.refresh(tag)
    return tag


def _create_product(db_session: Session, *, name: str = "KBC Lint") -> Product:
    product = Product(name=name)
    db_session.add(product)
    db_session.commit()
    db_session.refresh(product)
    return product


def _login(client: TestClient, email: str) -> None:
    client.post("/api/auth/login", json={"email": email, "password": "password123"})


def test_user_without_module_access_cannot_list_tags(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    _create_tagscan_module(db_session)
    _create_user(db_session, email="regular@example.com")
    _login(client, "regular@example.com")

    response = client.get("/api/modules/module-1/tags")

    assert response.status_code == 403


def test_user_with_access_but_no_tag_management_permission_cannot_list_tags(
    client: TestClient, db_session: Session
) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    user = _create_user(db_session, email="regular@example.com")
    _grant_module_access(db_session, user, module)
    _login(client, "regular@example.com")

    response = client.get("/api/modules/module-1/tags")

    assert response.status_code == 403


def test_tag_management_permission_does_not_grant_access_to_dashboard_screen(
    client: TestClient, db_session: Session
) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    user = _create_user(db_session, email="viewer@example.com")
    _grant_module_access(db_session, user, module)
    _grant_tag_management_permission(db_session, user, can_view=True, can_create=True)
    _login(client, "viewer@example.com")

    assert client.get("/api/modules/module-1/tags").status_code == 200
    assert client.get("/api/modules/module-1/files/tree").status_code == 403


def test_super_admin_can_register_a_tag_without_an_explicit_role(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, module)
    _login(client, "admin@example.com")

    response = client.post("/api/modules/module-1/tags", json={"epc_uid": "E200001122334455"})

    assert response.status_code == 201
    body = response.json()
    assert body["epc_uid"] == "E200001122334455"
    assert body["status"] == "active"


def test_view_only_tag_management_permission_cannot_create_a_tag(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    user = _create_user(db_session, email="viewer@example.com")
    _grant_module_access(db_session, user, module)
    _grant_tag_management_permission(db_session, user, can_view=True)
    _login(client, "viewer@example.com")

    response = client.post("/api/modules/module-1/tags", json={"epc_uid": "E200001122334455"})

    assert response.status_code == 403


def test_create_tag_with_full_field_set(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    tagscan_module = _create_tagscan_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, tagscan_module)
    product = _create_product(db_session, name="KBC Lint")
    _login(client, "admin@example.com")

    response = client.post(
        "/api/modules/module-1/tags",
        json={
            "epc_uid": "E200001122334455",
            "status": "lost",
            "assigned_product_id": product.id,
            "assigned_serial_number": "SN-001",
            "date_assigned": "2026-01-15",
            "last_read_at": "2026-02-01T10:30:00",
            "last_reader_id": "Reader-1",
            "last_location": "Warehouse A",
            "manufacturer": "Impinj",
            "batch_number": "B-42",
            "notes_1": "note 1",
            "notes_5": "note 5",
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "lost"
    assert body["assigned_product_id"] == product.id
    assert body["assigned_serial_number"] == "SN-001"
    assert body["manufacturer"] == "Impinj"
    assert body["notes_1"] == "note 1"
    assert body["notes_5"] == "note 5"


def test_creating_a_tag_with_an_unknown_assigned_product_id_returns_404(
    client: TestClient, db_session: Session
) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, module)
    _login(client, "admin@example.com")

    response = client.post(
        "/api/modules/module-1/tags", json={"epc_uid": "E200001122334455", "assigned_product_id": 999}
    )

    assert response.status_code == 404


def test_creating_a_tag_with_a_duplicate_epc_uid_is_rejected(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, module)
    _create_tag(db_session, epc_uid="E200001122334455")
    _login(client, "admin@example.com")

    response = client.post("/api/modules/module-1/tags", json={"epc_uid": "E200001122334455"})

    assert response.status_code == 409


def test_list_tags_ordered_by_epc_uid(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, module)
    _create_tag(db_session, epc_uid="E2ZZZ")
    _create_tag(db_session, epc_uid="E2AAA")
    _login(client, "admin@example.com")

    response = client.get("/api/modules/module-1/tags")

    assert response.status_code == 200
    assert [tag["epc_uid"] for tag in response.json()] == ["E2AAA", "E2ZZZ"]


def test_super_admin_can_update_a_tag(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, module)
    tag = _create_tag(db_session, epc_uid="E200001122334455")
    _login(client, "admin@example.com")

    response = client.put(
        f"/api/modules/module-1/tags/{tag.id}",
        json={"epc_uid": "E200001122334455", "status": "retired"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "retired"


def test_updating_a_missing_tag_returns_404(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, module)
    _login(client, "admin@example.com")

    response = client.put("/api/modules/module-1/tags/999", json={"epc_uid": "Nope"})

    assert response.status_code == 404


def test_super_admin_can_delete_a_tag(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, module)
    tag = _create_tag(db_session, epc_uid="E200001122334455")
    _login(client, "admin@example.com")

    response = client.delete(f"/api/modules/module-1/tags/{tag.id}")

    assert response.status_code == 204
    assert client.get("/api/modules/module-1/tags").json() == []


def test_deleting_a_missing_tag_returns_404(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, module)
    _login(client, "admin@example.com")

    response = client.delete("/api/modules/module-1/tags/999")

    assert response.status_code == 404


# --- XLSX template + bulk import --------------------------------------------

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _build_xlsx(header: list[str], rows: list[list]) -> bytes:
    """Build a minimal XLSX workbook in memory, for uploading in tests."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_downloading_the_template_requires_module_access(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    _create_tagscan_module(db_session)
    _create_user(db_session, email="regular@example.com")
    _login(client, "regular@example.com")

    response = client.get("/api/modules/module-1/tags/template")

    assert response.status_code == 403


def test_super_admin_can_download_the_template(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, module)
    _login(client, "admin@example.com")

    response = client.get("/api/modules/module-1/tags/template")

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_CONTENT_TYPE
    assert "attachment" in response.headers["content-disposition"]

    workbook = load_workbook(io.BytesIO(response.content))
    header_row = [cell.value for cell in workbook.active[1]]
    assert "epc_uid" in header_row
    assert "product_name" in header_row


def test_view_only_tag_management_permission_cannot_import_tags(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    user = _create_user(db_session, email="viewer@example.com")
    _grant_module_access(db_session, user, module)
    _grant_tag_management_permission(db_session, user, can_view=True)
    _login(client, "viewer@example.com")

    xlsx_bytes = _build_xlsx(["epc_uid"], [["E200001122334455"]])
    response = client.post(
        "/api/modules/module-1/tags/import",
        files={"file": ("tags.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)},
    )

    assert response.status_code == 403


def test_import_creates_updates_and_reports_row_errors(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, module)
    product = _create_product(db_session, name="KBC Lint")
    existing_tag = _create_tag(db_session, epc_uid="E2EXISTING")
    _login(client, "admin@example.com")

    xlsx_bytes = _build_xlsx(
        ["epc_uid", "status", "product_name", "manufacturer"],
        [
            ["E2NEW001", "active", "KBC Lint", "Impinj"],
            ["E2EXISTING", "retired", None, None],
            ["E2BADSTATUS", "not-a-status", None, None],
            ["E2NOPRODUCT", "active", "Nonexistent Product", None],
            [None, "active", None, None],
        ],
    )

    response = client.post(
        "/api/modules/module-1/tags/import",
        files={"file": ("tags.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)},
    )

    assert response.status_code == 200
    results = response.json()["results"]
    outcomes = {row["epc_uid"]: row["outcome"] for row in results}
    assert outcomes["E2NEW001"] == "created"
    assert outcomes["E2EXISTING"] == "updated"
    assert outcomes["E2BADSTATUS"] == "error"
    assert outcomes["E2NOPRODUCT"] == "error"
    missing_epc_row = next(row for row in results if row["row_number"] == 6)
    assert missing_epc_row["outcome"] == "error"

    # The valid rows actually persisted...
    new_tag = db_session.scalar(select(RfidTag).where(RfidTag.epc_uid == "E2NEW001"))
    assert new_tag is not None
    assert new_tag.assigned_product_id == product.id
    assert new_tag.manufacturer == "Impinj"

    db_session.refresh(existing_tag)
    assert existing_tag.status == "retired"

    # ...and the invalid rows didn't create anything.
    assert db_session.scalar(select(RfidTag).where(RfidTag.epc_uid == "E2BADSTATUS")) is None
    assert db_session.scalar(select(RfidTag).where(RfidTag.epc_uid == "E2NOPRODUCT")) is None


def test_import_accepts_native_excel_date_cells(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, module)
    _login(client, "admin@example.com")

    xlsx_bytes = _build_xlsx(
        ["epc_uid", "date_assigned", "last_read_at"],
        [["E2DATED", date(2026, 1, 15), datetime(2026, 2, 1, 10, 30)]],
    )

    response = client.post(
        "/api/modules/module-1/tags/import",
        files={"file": ("tags.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)},
    )

    assert response.status_code == 200
    assert response.json()["results"][0]["outcome"] == "created"

    tag = db_session.scalar(select(RfidTag).where(RfidTag.epc_uid == "E2DATED"))
    assert tag is not None
    assert tag.date_assigned == date(2026, 1, 15)
    assert tag.last_read_at.replace(tzinfo=None) == datetime(2026, 2, 1, 10, 30)
