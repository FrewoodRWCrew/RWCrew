# These tests check the Scanners tile of TagScan's "Data Upload/Download"
# screen: template, XLSX import (create / update by name / row errors) and
# export, gated by the "tagscan.dataupload" screen permission.

import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models.scanner import Scanner
from app.modules.module_1.screens import sync_screens as sync_tagscan_screens
from tests.modules.module_1.test_scanners import (
    _create_product_type,
    _create_scanner,
    _create_tagscan_module,
    _create_user,
    _grant_module_access,
    _grant_screen_permission,
    _login,
)

BASE = "/api/modules/module-1"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _build_xlsx(header: list[str], rows: list[list]) -> bytes:
    """Build a minimal XLSX workbook in memory, for uploading in tests."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _admin(client: TestClient, db_session: Session) -> None:
    """Log in as a super admin with TagScan module access."""
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    admin = _create_user(db_session, email="admin@example.com", is_super_admin=True)
    _grant_module_access(db_session, admin, module)
    _login(client, "admin@example.com")


def _upload(client: TestClient, xlsx_bytes: bytes):
    return client.post(f"{BASE}/scanner-import", files={"file": ("scanners.xlsx", xlsx_bytes, XLSX_CONTENT_TYPE)})


def test_template_has_the_scanner_columns(client: TestClient, db_session: Session) -> None:
    _admin(client, db_session)

    response = client.get(f"{BASE}/scanner-import/template")

    assert response.status_code == 200
    header = [cell.value for cell in load_workbook(io.BytesIO(response.content)).active[1]]
    assert header[:3] == ["scanner", "type_name", "technology"]


def test_import_creates_updates_and_reports_row_errors(client: TestClient, db_session: Session) -> None:
    _admin(client, db_session)
    product_type = _create_product_type(db_session, name="Reader Hardware")
    existing = _create_scanner(db_session, scanner="Scan_01", type_id=product_type.id)
    existing.api_key_hash = "kept-hash"
    db_session.commit()

    response = _upload(
        client,
        _build_xlsx(
            ["scanner", "type_name", "technology", "location"],
            [
                ["Scan_02", "reader hardware", "raspberry pi 5", "Gate B"],
                [" Scan_01 ", "Reader Hardware", "Other", "Dock"],
                ["Scan_03", "Unknown Type", "Other", None],
                ["Scan_04", "Reader Hardware", "Arduino", None],
                [None, "Reader Hardware", "Other", None],
            ],
        ),
    )

    assert response.status_code == 200
    results = response.json()["results"]
    assert [(row["scanner"], row["outcome"]) for row in results] == [
        ("Scan_02", "created"),
        ("Scan_01", "updated"),
        ("Scan_03", "error"),
        ("Scan_04", "error"),
        (None, "error"),
    ]

    created = db_session.scalar(select(Scanner).where(Scanner.scanner == "Scan_02"))
    assert (created.type_id, created.technology, created.location) == (product_type.id, "Raspberry Pi 5", "Gate B")
    db_session.refresh(existing)
    # Updated from the sheet, while its device API key stays untouched.
    assert (existing.technology, existing.location, existing.api_key_hash) == ("Other", "Dock", "kept-hash")
    assert db_session.scalar(select(Scanner).where(Scanner.scanner.in_(["Scan_03", "Scan_04"]))) is None


def test_export_lists_every_scanner_with_its_type_name(client: TestClient, db_session: Session) -> None:
    _admin(client, db_session)
    product_type = _create_product_type(db_session, name="Reader Hardware")
    _create_scanner(db_session, scanner="Scan_01", type_id=product_type.id, technology="Raspberry Pi 4")

    response = client.get(f"{BASE}/scanners/export")

    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert [cell.value for cell in sheet[2]][:3] == ["Scan_01", "Reader Hardware", "Raspberry Pi 4"]


def test_view_only_dataupload_permission_cannot_import_scanners(client: TestClient, db_session: Session) -> None:
    sync_tagscan_screens(db_session)
    module = _create_tagscan_module(db_session)
    user = _create_user(db_session, email="viewer@example.com")
    _grant_module_access(db_session, user, module)
    _grant_screen_permission(db_session, user, "tagscan.dataupload", can_view=True)
    _login(client, "viewer@example.com")

    # View is enough for the export, not for the upload.
    assert client.get(f"{BASE}/scanners/export").status_code == 200
    response = _upload(client, _build_xlsx(["scanner"], [["Scan_09"]]))
    assert response.status_code == 403
