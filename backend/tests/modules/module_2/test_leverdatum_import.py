# These tests check the Leverdata tile on KarTracker's "Data Upload/
# Download" screen: the pre-filled XLSX template, the bulk import (an upsert
# per festival, rows without dates skipped) and the export — all behind the
# same "kartracker.dataupload" screen key as the other tiles.

import io
from datetime import date, datetime

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models.festival import Festival
from app.db.models.kartracker_leverdatum import KarTrackerLeverdatum
from app.db.models.kartracker_user_role import KarTrackerUserRole
from app.db.models.season import Season
from app.modules.module_2.screens import sync_screens
from tests.modules.module_2.conftest import (
    create_kartracker_module,
    create_role_with_permissions,
    create_user,
    grant_module_access,
    login,
)

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER = ["season_name", "festival_name", "delivery_date", "pickup_date"]


def _build_xlsx(rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADER)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _read_rows(content: bytes) -> list[tuple]:
    sheet = load_workbook(io.BytesIO(content)).active
    return list(sheet.iter_rows(values_only=True))


def _login_as_admin(client: TestClient, db_session: Session) -> None:
    sync_screens(db_session)
    module = create_kartracker_module(db_session)
    admin = create_user(db_session, email="admin@example.com", is_super_admin=True)
    grant_module_access(db_session, admin, module)
    login(client, "admin@example.com")


def _create_season(db_session: Session, name: str = "2026") -> Season:
    season = Season(name=name)
    db_session.add(season)
    db_session.commit()
    db_session.refresh(season)
    return season


def _create_festival(db_session: Session, season: Season, name: str, *, active: bool = True) -> Festival:
    festival = Festival(
        name=name, start_date=date(2026, 7, 1), end_date=date(2026, 7, 3), season_id=season.id, active=active
    )
    db_session.add(festival)
    db_session.commit()
    db_session.refresh(festival)
    return festival


def _upload(client: TestClient, rows: list[list]):
    return client.post(
        "/api/modules/module-2/leverdatum-import",
        files={"file": ("leverdata.xlsx", _build_xlsx(rows), XLSX_CONTENT_TYPE)},
    )


# --- template + export -------------------------------------------------------


def test_template_lists_active_festivals_with_empty_dates(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    season = _create_season(db_session)
    _create_festival(db_session, season, "Summer Bash")
    _create_festival(db_session, season, "Old Fest", active=False)

    response = client.get("/api/modules/module-2/leverdatum-import/template")

    assert response.status_code == 200
    rows = _read_rows(response.content)
    assert rows[0] == tuple(HEADER)
    assert rows[1:] == [("2026", "Summer Bash", None, None)]


def test_export_contains_saved_dates(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    season = _create_season(db_session)
    festival = _create_festival(db_session, season, "Summer Bash")
    db_session.add(
        KarTrackerLeverdatum(festival_id=festival.id, delivery_date=date(2026, 6, 28), pickup_date=date(2026, 7, 5))
    )
    db_session.commit()

    response = client.get("/api/modules/module-2/leverdata/export")

    assert response.status_code == 200
    rows = _read_rows(response.content)
    assert rows[1] == ("2026", "Summer Bash", datetime(2026, 6, 28), datetime(2026, 7, 5))


# --- bulk import ---------------------------------------------------------------


def test_view_only_dataupload_permission_cannot_import(client: TestClient, db_session: Session) -> None:
    sync_screens(db_session)
    module = create_kartracker_module(db_session)
    user = create_user(db_session, email="viewer@example.com")
    grant_module_access(db_session, user, module)
    role = create_role_with_permissions(db_session, name="Viewer", screen_key="kartracker.dataupload", can_view=True)
    db_session.add(KarTrackerUserRole(user_id=user.id, role_id=role.id))
    db_session.commit()
    login(client, "viewer@example.com")

    assert client.get("/api/modules/module-2/leverdatum-import/template").status_code == 200
    assert _upload(client, []).status_code == 403


def test_import_creates_updates_skips_and_reports_errors(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    season = _create_season(db_session)
    new_festival = _create_festival(db_session, season, "Summer Bash")
    existing_festival = _create_festival(db_session, season, "Winter Fest")
    _create_festival(db_session, season, "Empty Fest")
    db_session.add(KarTrackerLeverdatum(festival_id=existing_festival.id, delivery_date=date(2026, 1, 1)))
    db_session.commit()

    response = _upload(
        client,
        [
            ["2026", "summer bash", datetime(2026, 6, 28), "05/07/2026"],  # created (case-insensitive, text date)
            ["2026", "Winter Fest", "2026-12-01", None],  # updated
            ["2026", "Empty Fest", None, None],  # skipped
            ["2026", "Nope", "2026-06-01", None],  # unknown festival
            ["1999", "Summer Bash", "2026-06-01", None],  # unknown season
            ["2026", "Summer Bash", "2026-06-01", None],  # same festival twice
        ],
    )

    assert response.status_code == 200
    outcomes = [(row["row_number"], row["outcome"]) for row in response.json()["results"]]
    assert outcomes == [(2, "created"), (3, "updated"), (4, "skipped"), (5, "error"), (6, "error"), (7, "error")]

    db_session.expire_all()
    created = db_session.scalar(select(KarTrackerLeverdatum).where(KarTrackerLeverdatum.festival_id == new_festival.id))
    assert (created.delivery_date, created.pickup_date) == (date(2026, 6, 28), date(2026, 7, 5))
    updated = db_session.scalar(
        select(KarTrackerLeverdatum).where(KarTrackerLeverdatum.festival_id == existing_festival.id)
    )
    assert (updated.delivery_date, updated.pickup_date) == (date(2026, 12, 1), None)


def test_import_rejects_bad_dates_and_pickup_before_delivery(client: TestClient, db_session: Session) -> None:
    _login_as_admin(client, db_session)
    season = _create_season(db_session)
    _create_festival(db_session, season, "Summer Bash")
    _create_festival(db_session, season, "Winter Fest")

    response = _upload(
        client,
        [
            ["2026", "Summer Bash", "not a date", None],
            ["2026", "Winter Fest", "2026-07-05", "2026-07-01"],
        ],
    )

    assert response.status_code == 200
    assert [row["outcome"] for row in response.json()["results"]] == ["error", "error"]
    assert db_session.scalar(select(KarTrackerLeverdatum)) is None
